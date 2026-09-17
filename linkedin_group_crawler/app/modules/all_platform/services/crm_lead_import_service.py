from __future__ import annotations

import io
import re
import unicodedata
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.core.config import settings
from app.core.supabase_client import execute_supabase_query, get_supabase_client
from app.modules.all_platform.services.crm_lead_service import DuplicateLeadError, create_lead, is_valid_email
from app.modules.all_platform.services.crm_customer_service import normalize_email, normalize_phone
from app.modules.all_platform.services.crm_permission_service import can_manage_shared_master_data, has_full_crm_access
from app.modules.all_platform.services.supabase_categories_service import add_category

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 1000

FIELD_HEADERS = {
    "lead_name": "Họ tên người liên hệ",
    "company_name": "Công ty/Tổ chức",
    "phone": "SĐT",
    "email": "Email",
    "position": "Chức vụ",
    "source": "Nguồn",
    "owner": "Người phụ trách Lead",
    "zalo": "Zalo",
    "facebook": "Facebook",
    "telegram": "Telegram",
    "website": "Website",
    "note": "Ghi chú",
}


def _key(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or "").strip().lower())
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", "", text.replace("đ", "d"))


HEADER_ALIASES = {
    _key(label): field for field, label in FIELD_HEADERS.items()
}
HEADER_ALIASES.update({
    "hoten": "lead_name", "tenlead": "lead_name", "fullname": "lead_name",
    "congty": "company_name", "tochuc": "company_name", "company": "company_name",
    "sodienthoai": "phone", "phone": "phone", "mobile": "phone",
    "emailaddress": "email", "position": "position", "jobtitle": "position",
    "source": "source", "nguoiphutrach": "owner", "owner": "owner", "sdr": "owner",
    "notes": "note",
})


def _tenant_categories(category_type: str) -> list[dict[str, Any]]:
    # Historical helper name: categories are shared CRM master data (migration
    # 124), not tenant-scoped entities. Match the active options on Lead forms;
    # only Lead read/write/dedup queries must filter settings.crm_instance.
    supabase = get_supabase_client()
    res = execute_supabase_query(
        lambda: supabase.table("categories")
        .select("id, code, name, category_type, is_active")
        .eq("category_type", category_type)
        .eq("is_active", True)
        .execute()
    )
    return res.data or []


def _active_users() -> list[dict[str, Any]]:
    supabase = get_supabase_client()
    res = execute_supabase_query(
        lambda: supabase.table("app_users")
        .select("id, name, email, is_active")
        .eq("is_active", True)
        .execute()
    )
    return res.data or []


def _lookup(rows: list[dict[str, Any]], fields: tuple[str, ...]) -> dict[str, list[dict[str, Any]]]:
    result: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        for field in fields:
            value = row.get(field)
            if value:
                result.setdefault(_key(value), []).append(row)
    return result


def _excel_text(value: Any) -> Any:
    """Prevent CRM master labels from being interpreted as formulas in the
    downloadable workbook."""
    if not isinstance(value, str):
        return value
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def build_template() -> bytes:
    sources = _tenant_categories("crm_source")
    positions = _tenant_categories("crm_position")
    users = _active_users()
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    headers = list(FIELD_HEADERS.values())
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    widths = [28, 28, 18, 28, 22, 20, 28, 18, 24, 18, 24, 36]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + index)].width = width
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2, max_row=1001, min_col=3, max_col=3):
        row[0].number_format = "@"

    lookup = wb.create_sheet("Danh mục tham chiếu")
    lookup.append(["Nguồn", "Mã nguồn", "Chức vụ", "Người phụ trách", "Email người phụ trách"])
    longest = max(len(sources), len(positions), len(users), 1)
    for index in range(longest):
        source = sources[index] if index < len(sources) else {}
        position = positions[index] if index < len(positions) else {}
        owner = users[index] if index < len(users) else {}
        lookup.append([
            _excel_text(source.get("name") or source.get("code")), _excel_text(source.get("code")),
            _excel_text(position.get("name") or position.get("code")), _excel_text(owner.get("name")), _excel_text(owner.get("email")),
        ])
    lookup.sheet_state = "hidden"
    if sources:
        dv = DataValidation(type="list", formula1=f'=INDIRECT("\'Danh mục tham chiếu\'!$A$2:$A${len(sources)+1}")')
        ws.add_data_validation(dv); dv.add("F2:F1001")
    if positions:
        dv = DataValidation(type="list", formula1=f'=INDIRECT("\'Danh mục tham chiếu\'!$C$2:$C${len(positions)+1}")')
        ws.add_data_validation(dv); dv.add("E2:E1001")
    if users:
        dv = DataValidation(type="list", formula1=f'=INDIRECT("\'Danh mục tham chiếu\'!$D$2:$D${len(users)+1}")')
        ws.add_data_validation(dv); dv.add("G2:G1001")
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _read_rows(raw: bytes) -> tuple[list[str], list[tuple[int, dict[str, Any]]], dict[str, str]]:
    if len(raw) > MAX_IMPORT_BYTES:
        raise ValueError("File vuot qua gioi han 5 MB.")
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Khong doc duoc file Excel .xlsx.") from exc
    ws = wb.active
    values = ws.iter_rows(values_only=True)
    raw_headers = next(values, None)
    if not raw_headers:
        raise ValueError("File Excel khong co dong tieu de.")
    headers = [str(value or "").strip() for value in raw_headers]
    mapping: dict[str, str] = {}
    for index, header in enumerate(headers):
        field = HEADER_ALIASES.get(_key(header))
        if field and field not in mapping:
            mapping[field] = str(index)
    if "lead_name" not in mapping:
        raise ValueError("Khong tim thay cot Ho ten nguoi lien he.")
    parsed: list[tuple[int, dict[str, Any]]] = []
    for row_number, row in enumerate(values, 2):
        if not any(value not in (None, "") for value in row):
            continue
        if len(parsed) >= MAX_IMPORT_ROWS:
            raise ValueError(f"File vuot qua gioi han {MAX_IMPORT_ROWS} dong du lieu.")
        data: dict[str, Any] = {}
        for field, raw_index in mapping.items():
            index = int(raw_index)
            value = row[index] if index < len(row) else None
            data[field] = str(value).strip() if value is not None else None
        parsed.append((row_number, data))
    return headers, parsed, {field: headers[int(index)] for field, index in mapping.items()}


def _issue(column: str, code: str, message: str, **extra: Any) -> dict[str, Any]:
    payload: dict[str, Any] = {"column": FIELD_HEADERS.get(column, column), "code": code, "message": message}
    payload.update(extra)
    return payload


def _validate_and_resolve_rows(
    parsed: list[tuple[int, dict[str, Any]]],
    user: dict[str, Any],
    *,
    allow_create_master: bool,
) -> dict[str, Any]:
    sources = _lookup(_tenant_categories("crm_source"), ("code", "name"))
    positions = _lookup(_tenant_categories("crm_position"), ("id", "code", "name"))
    users = _lookup(_active_users(), ("id", "email", "name"))
    supabase = get_supabase_client()
    incoming_phones = {normalize_phone(row.get("phone")) for _, row in parsed}
    incoming_emails = {normalize_email(row.get("email")) for _, row in parsed}
    incoming_phones.discard(None)
    incoming_emails.discard(None)
    existing_by_id: dict[str, dict[str, Any]] = {}
    if incoming_phones:
        phone_res = execute_supabase_query(
            lambda: supabase.table("crm_leads")
            .select("id, lead_name, phone, phone_normalized, email_normalized")
            .eq("instance", settings.crm_instance)
            .in_("phone_normalized", list(incoming_phones))
            .execute()
        )
        existing_by_id.update({str(row["id"]): row for row in phone_res.data or []})
        # Historical rows may predate phone_normalized. Scan only that legacy
        # subset, not the entire tenant Lead table.
        legacy_res = execute_supabase_query(
            lambda: supabase.table("crm_leads")
            .select("id, lead_name, phone, phone_normalized, email_normalized")
            .eq("instance", settings.crm_instance)
            .is_("phone_normalized", "null")
            .not_.is_("phone", "null")
            .execute()
        )
        for row in legacy_res.data or []:
            if normalize_phone(row.get("phone")) in incoming_phones:
                existing_by_id[str(row["id"])] = row
    if incoming_emails:
        email_res = execute_supabase_query(
            lambda: supabase.table("crm_leads")
            .select("id, lead_name, phone, phone_normalized, email_normalized")
            .eq("instance", settings.crm_instance)
            .in_("email_normalized", list(incoming_emails))
            .execute()
        )
        existing_by_id.update({str(row["id"]): row for row in email_res.data or []})
    existing_phones: dict[str, set[str]] = {}
    existing_emails: dict[str, set[str]] = {}
    for existing in existing_by_id.values():
        existing_id = str(existing.get("id") or "")
        stored_phone = existing.get("phone_normalized") or normalize_phone(existing.get("phone"))
        stored_email = existing.get("email_normalized")
        if stored_phone: existing_phones.setdefault(str(stored_phone), set()).add(existing_id)
        if stored_email: existing_emails.setdefault(str(stored_email), set()).add(existing_id)
    actor_id = str(user.get("id") or "")
    can_assign = has_full_crm_access(user)
    seen_phone: dict[str, int] = {}
    seen_email: dict[str, int] = {}
    result_rows: list[dict[str, Any]] = []

    for row_number, source_row in parsed:
        row = dict(source_row)
        issues: list[dict[str, Any]] = []
        pending_creates: list[dict[str, str]] = []
        name = str(row.get("lead_name") or "").strip()
        phone = str(row.get("phone") or "").strip()
        email = str(row.get("email") or "").strip()
        phone_norm = normalize_phone(phone)
        email_norm = normalize_email(email)
        if not name:
            issues.append(_issue("lead_name", "required", "Bắt buộc nhập họ tên."))
        if not phone and not email:
            issues.append(_issue("phone", "required_contact", "Cần có SĐT hoặc Email."))
        if phone and not phone_norm:
            issues.append(_issue("phone", "invalid_phone", "Số điện thoại không đúng định dạng."))
        if email and not is_valid_email(email):
            issues.append(_issue("email", "invalid_email", "Email không đúng định dạng."))

        source_value = str(row.get("source") or "").strip()
        if source_value:
            choices = sources.get(_key(source_value), [])
            unique = list({item["id"]: item for item in choices}.values())
            if len(unique) == 1:
                row["source"] = unique[0].get("code")
            elif len(unique) == 0 and allow_create_master:
                pending_creates.append({"field": "source", "value": source_value})
            elif len(unique) == 0:
                issues.append(_issue(
                    "source", "unknown_source",
                    f"Nguồn '{source_value}' chưa tồn tại trong CRM và bạn không có quyền tạo mới.",
                ))
            else:
                issues.append(_issue(
                    "source", "unknown_source",
                    f"Nguồn '{source_value}' khớp nhiều hơn 1 danh mục, vui lòng chọn lại.",
                ))

        position_value = str(row.get("position") or "").strip()
        if position_value:
            choices = positions.get(_key(position_value), [])
            unique = list({item["id"]: item for item in choices}.values())
            if len(unique) == 1:
                row["position_category_id"] = unique[0]["id"]
            elif len(unique) == 0 and allow_create_master:
                pending_creates.append({"field": "position", "value": position_value})
            elif len(unique) == 0:
                issues.append(_issue(
                    "position", "unknown_position",
                    f"Chức vụ '{position_value}' chưa tồn tại trong CRM và bạn không có quyền tạo mới.",
                ))
            else:
                issues.append(_issue(
                    "position", "unknown_position",
                    f"Chức vụ '{position_value}' khớp nhiều hơn 1 danh mục, vui lòng chọn lại.",
                ))

        # "owner" always holds the CURRENT display text/id (raw Excel text on
        # first parse, a member id string once Sale picks from the FE member
        # combobox) and is re-resolved fresh on every pass, exactly like
        # source/position above — never popped/discarded. Popping it after a
        # failed match (earlier version of this code) silently lost the
        # unresolved text on the next revalidate round-trip and made the row
        # default back to "assign to me", masking a real unknown_owner error.
        # `users` indexes by id as well as email/name, so a member id string
        # resolves through the exact same lookup path as typed name/email.
        owner_value = str(row.get("owner") or "").strip()
        if owner_value:
            choices = users.get(_key(owner_value), [])
            unique = list({item["id"]: item for item in choices}.values())
            if len(unique) == 1:
                row["sdr_id"] = str(unique[0]["id"])
                row["owner_name"] = unique[0].get("name") or unique[0].get("email")
                if not can_assign and str(unique[0]["id"]) != actor_id:
                    issues.append(_issue(
                        "owner", "owner_forbidden",
                        f"Bạn không có quyền gán Lead cho {unique[0].get('name') or unique[0].get('email')}.",
                    ))
            else:
                issues.append(_issue(
                    "owner", "unknown_owner",
                    f"Không tìm thấy đúng 1 thành viên CRM khớp với '{owner_value}'.",
                ))
        else:
            row["sdr_id"] = actor_id or None
            row["owner_name"] = user.get("name") or user.get("email") or None

        phone_matches = existing_phones.get(phone_norm or "", set())
        email_matches = existing_emails.get(email_norm or "", set())
        duplicate_matches = phone_matches | email_matches
        duplicate_row = None
        if phone_norm and phone_norm in seen_phone:
            duplicate_row = seen_phone[phone_norm]
        if email_norm and email_norm in seen_email:
            duplicate_row = duplicate_row or seen_email[email_norm]
        if phone_norm: seen_phone.setdefault(phone_norm, row_number)
        if email_norm: seen_email.setdefault(email_norm, row_number)

        identity_conflict = bool(phone_matches and email_matches and phone_matches.isdisjoint(email_matches))
        if identity_conflict:
            issues.append(_issue("email", "identity_conflict", "SĐT và Email đang thuộc hai Lead khác nhau trong hệ thống."))
        status = "error" if issues else "duplicate" if duplicate_matches or duplicate_row else "valid"
        if duplicate_matches and not identity_conflict:
            match_id = next(iter(duplicate_matches))
            match_row = existing_by_id.get(match_id, {})
            match_name = match_row.get("lead_name") or "Lead khác"
            issues.append(_issue(
                "phone", "duplicate_database",
                f"Trùng SĐT/Email với Lead có sẵn — {match_name}.",
                lead_id=match_id, lead_name=match_row.get("lead_name"),
            ))
        elif duplicate_row:
            issues.append(_issue(
                "phone", "duplicate_file",
                f"Trùng SĐT/Email với dòng {duplicate_row} trong file.",
                duplicate_row_number=duplicate_row,
            ))
        result_rows.append({
            "row_number": row_number,
            "status": status,
            "data": row,
            "issues": issues,
            "pending_creates": pending_creates,
        })

    summary = {
        "total": len(result_rows),
        "valid": sum(row["status"] == "valid" for row in result_rows),
        "duplicate": sum(row["status"] == "duplicate" for row in result_rows),
        "error": sum(row["status"] == "error" for row in result_rows),
    }
    return {"summary": summary, "rows": result_rows}


def preview_import(raw: bytes, user: dict[str, Any]) -> dict[str, Any]:
    headers, parsed, mapping = _read_rows(raw)
    result = _validate_and_resolve_rows(parsed, user, allow_create_master=can_manage_shared_master_data(user))
    result["headers"] = headers
    result["mapping"] = mapping
    result["can_create_master"] = can_manage_shared_master_data(user)
    return result


def revalidate_rows(rows_payload: list[dict[str, Any]], user: dict[str, Any]) -> dict[str, Any]:
    """Re-run the same validation core on rows already edited client-side
    (no Excel file involved). Always revalidates the FULL row set sent, not
    just the row that was edited, so in-file/in-batch duplicate detection
    (seen_phone/seen_email) stays correct."""
    parsed = [
        (int(row.get("row_number") or 0), {key: value for key, value in row.items() if key != "row_number"})
        for row in rows_payload
    ]
    result = _validate_and_resolve_rows(parsed, user, allow_create_master=can_manage_shared_master_data(user))
    result["can_create_master"] = can_manage_shared_master_data(user)
    return result


def _resolve_or_create_pending_masters(rows: list[dict[str, Any]], user: dict[str, Any]) -> None:
    """Mutates `rows` in place: for every pending_creates entry on a row,
    reuse an existing category if one now matches (re-checked fresh, in case
    another row/request created it since preview), otherwise create exactly
    one new category per distinct new value across the whole batch — never
    one per row. No-op (silently drops pending_creates) if the user lost
    master-data permission between preview and confirm."""
    if not can_manage_shared_master_data(user):
        for row in rows:
            row["pending_creates"] = []
        return
    type_map = {"source": "crm_source", "position": "crm_position"}
    values_by_field: dict[str, dict[str, str]] = {"source": {}, "position": {}}
    for row in rows:
        for pending in row.get("pending_creates") or []:
            field = pending.get("field")
            value = str(pending.get("value") or "").strip()
            if field in values_by_field and value:
                values_by_field[field].setdefault(_key(value), value)

    resolved: dict[str, dict[str, dict[str, Any]]] = {"source": {}, "position": {}}
    for field, values_by_key in values_by_field.items():
        if not values_by_key:
            continue
        category_type = type_map[field]
        fields = ("code", "name") if field == "source" else ("id", "code", "name")
        existing = _lookup(_tenant_categories(category_type), fields)
        for key, raw_value in values_by_key.items():
            matches = existing.get(key)
            if matches:
                resolved[field][key] = next(iter({item["id"]: item for item in matches}.values()))
            else:
                resolved[field][key] = add_category({
                    "category_type": category_type,
                    "code": raw_value,
                    "name": raw_value,
                })

    for row in rows:
        for pending in row.get("pending_creates") or []:
            field = pending.get("field")
            value = str(pending.get("value") or "").strip()
            created_row = resolved.get(field, {}).get(_key(value))
            if not created_row:
                continue
            if field == "source":
                row["data"]["source"] = created_row.get("code")
            elif field == "position":
                row["data"]["position_category_id"] = created_row.get("id")
        row["pending_creates"] = []


def confirm_import(rows_payload: list[dict[str, Any]], selected_rows: list[int], user: dict[str, Any]) -> dict[str, Any]:
    if not selected_rows:
        raise ValueError("Vui lòng chọn ít nhất một dòng hợp lệ để import.")
    revalidated = revalidate_rows(rows_payload, user)
    by_number = {row["row_number"]: row for row in revalidated["rows"]}
    requested = sorted(set(selected_rows))
    pending_targets = [by_number[number] for number in requested if by_number.get(number, {}).get("status") == "valid"]
    _resolve_or_create_pending_masters(pending_targets, user)

    created: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    for row_number in requested:
        row = by_number.get(row_number)
        if not row or row["status"] != "valid":
            skipped.append({"row_number": row_number, "message": "Dòng không còn hợp lệ sau khi kiểm tra lại."})
            continue
        # "owner"/"owner_name" are display-only — crm_leads has no such
        # columns, only "sdr_id" (already resolved above).
        data = {key: value for key, value in row["data"].items() if key not in ("owner", "owner_name")}
        data["status"] = "mql"
        try:
            created.append(create_lead(data, user))
        except DuplicateLeadError:
            skipped.append({"row_number": row_number, "message": "Lead vừa bị trùng SĐT/Email."})
        except Exception as exc:
            skipped.append({"row_number": row_number, "message": str(exc)})
    return {
        "created": len(created),
        "skipped": len(skipped),
        "requested": len(requested),
        "created_lead_ids": [row.get("id") for row in created],
        "errors": skipped,
    }
