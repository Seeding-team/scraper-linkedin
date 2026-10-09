"""Document engine cho AI Contract Copilot: chỉnh sửa TRỰC TIẾP trên bản sao DOCX của mẫu, xuất PDF từ chính DOCX đó.

Nguyên tắc (xem docs/CONTRACT_DOCX_PIPELINE_2026-10-09.md):
  * File mẫu gốc không bao giờ bị ghi đè - mọi thao tác chạy trên bytes được nạp lại thành Document mới.
  * AI chỉ ĐỀ XUẤT (id đoạn -> text mới). Engine này là nơi DUY NHẤT áp thay đổi vào tài liệu, và chỉ đổi text của run
    (giữ nguyên style, numbering, section, header/footer, bảng, page break, khu vực chữ ký).
  * Số liệu tiền/VAT chỉ đến từ dữ liệu CRM (báo giá) - không do AI sinh ra.
  * PDF = chuyển đổi từ DOCX đã chỉnh bằng LibreOffice headless; thiếu engine thì BÁO LỖI, không thay bằng cách khác.
"""
from __future__ import annotations

import copy
import io
import os
import re
import shutil
import signal
import subprocess
import tempfile
import threading
import time
import zipfile
from dataclasses import dataclass, field
from typing import Any, Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


class PdfEngineUnavailable(RuntimeError):
    """Server không có LibreOffice (soffice) hoặc chuyển đổi thất bại - KHÔNG được thay bằng bản PDF dựng cách khác."""


# ───────────────────────── đọc cấu trúc ─────────────────────────

def load_document(content: bytes) -> DocumentType:
    validate_docx_bytes(content)
    return Document(io.BytesIO(content))


def _all_paragraph_elements(doc: DocumentType) -> list[Any]:
    """Mọi <w:p> trong thân tài liệu theo thứ tự xuất hiện (gồm cả đoạn trong bảng, bảng lồng). Header/footer KHÔNG nằm trong đây."""
    return list(doc.element.body.iter(qn("w:p")))


def paragraph_by_id(doc: DocumentType, pid: str) -> Paragraph | None:
    m = re.fullmatch(r"p(\d+)", pid or "")
    if not m:
        return None
    els = _all_paragraph_elements(doc)
    idx = int(m.group(1))
    return Paragraph(els[idx], doc._body) if idx < len(els) else None


def _in_table(p_el: Any) -> bool:
    return any(a.tag == qn("w:tc") for a in p_el.iterancestors())


def _para_text(p_el: Any) -> str:
    return "".join(t.text or "" for t in p_el.iter(qn("w:t")))


@dataclass
class ParaInfo:
    id: str
    text: str
    style: str
    in_table: bool
    numbered: bool


def describe_paragraphs(doc: DocumentType) -> list[ParaInfo]:
    out: list[ParaInfo] = []
    for i, el in enumerate(_all_paragraph_elements(doc)):
        text = _para_text(el)
        ppr = el.find(qn("w:pPr"))
        style = ""
        numbered = False
        if ppr is not None:
            st = ppr.find(qn("w:pStyle"))
            style = st.get(qn("w:val")) if st is not None else ""
            numbered = ppr.find(qn("w:numPr")) is not None
        out.append(ParaInfo(f"p{i}", text, style or "", _in_table(el), numbered))
    return out


def fingerprint(doc: DocumentType) -> dict[str, Any]:
    """Dấu vân tay cấu trúc để kiểm tra 'giữ form': số đoạn/bảng/section, kích thước & lề trang, header/footer, page break."""
    body = doc.element.body
    sections = []
    for s in doc.sections:
        sections.append({
            "page": (s.page_width, s.page_height), "margins": (s.left_margin, s.right_margin, s.top_margin, s.bottom_margin),
            "header": "\n".join(p.text for p in s.header.paragraphs), "footer": "\n".join(p.text for p in s.footer.paragraphs),
        })
    tables = [(len(t.rows), len(t.columns)) for t in doc.tables]
    return {
        "paragraphs": sum(1 for p in _all_paragraph_elements(doc) if not _in_table(p)),   # đoạn trong bảng nằm ở so sánh bảng
        "tables": tables,
        "sections": sections,
        "pageBreaks": len(list(body.iter(qn("w:br")))) and sum(1 for b in body.iter(qn("w:br")) if b.get(qn("w:type")) == "page"),
        "sectPr": len(list(body.iter(qn("w:sectPr")))),
        "numberedParagraphs": sum(1 for p in describe_paragraphs(doc) if p.numbered and not p.in_table),
        "styles": sorted({p.style for p in describe_paragraphs(doc) if p.style and not p.in_table}),
    }


# ───────────────────────── thay text giữ định dạng ─────────────────────────

def _text_runs(p_el: Any) -> list[Any]:
    """<w:r> trực tiếp (cả trong <w:hyperlink>) có chứa <w:t>."""
    return [r for r in p_el.iter(qn("w:r")) if r.find(qn("w:t")) is not None]


def _set_run_text(run_el: Any, text: str) -> None:
    ts = run_el.findall(qn("w:t"))
    for extra in ts[1:]:
        run_el.remove(extra)
    t = ts[0]
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _run_text(r: Any) -> str:
    return "".join(t.text or "" for t in r.findall(qn("w:t")))


def _set_run_bold(run_el: Any, bold: bool) -> None:
    from docx.oxml import OxmlElement

    rpr = run_el.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        run_el.insert(0, rpr)
    for tag in ("w:b", "w:bCs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), "1" if bold else "0")


def _unbold_label_suffix(p_el: Any, label: str) -> None:
    """Đảm bảo phần văn bản TỪ `label` trở đi (vd 'Chức vụ: ...') KHÔNG BAO GIỜ in đậm, dù run chứa nó đang gộp chung
    với phần TRƯỚC đang in đậm (vd tên người đại diện) trong CÙNG 1 run - tách run thành 2 nếu cần (giữ nguyên định
    dạng phần trước, ép phần nhãn trở đi về không đậm) thay vì chỉ đổi text mà giữ nguyên rPr cũ của run gốc."""
    for r in _text_runs(p_el):
        t = _run_text(r)
        idx = t.find(label)
        if idx == -1:
            continue
        if idx == 0:
            _set_run_bold(r, False)
            continue
        t_el = r.find(qn("w:t"))
        before, after = t[:idx], t[idx:]
        t_el.text = before
        new_r = copy.deepcopy(r)
        new_r.find(qn("w:t")).text = after
        r.addnext(new_r)
        _set_run_bold(new_r, False)


def replace_range(p_el: Any, start: int, end: int, new: str) -> None:
    """Thay ký tự [start,end) của text đoạn bằng `new`; chữ ngoài vùng đổi giữ nguyên run/format của nó.
    Run đầu tiên chạm vùng đổi nhận `new` (kế thừa định dạng của nó); các run chạm vùng đổi khác chỉ bị cắt phần trùng."""
    runs = _text_runs(p_el)
    if not runs:
        return
    spans, pos = [], 0
    for r in runs:
        t = _run_text(r)
        spans.append((r, t, pos, pos + len(t)))
        pos += len(t)
    if start == end:                                          # chèn thuần
        for r, t, rs, re_ in spans:
            if rs <= start <= re_ and (start < re_ or r is runs[-1]):
                k = start - rs
                _set_run_text(r, t[:k] + new + t[k:])
                return
        r, t, _, _ = spans[-1]
        _set_run_text(r, t + new)
        return
    placed = False
    for r, t, rs, re_ in spans:
        if re_ <= start or rs >= end:
            continue
        lo, hi = max(start, rs) - rs, min(end, re_) - rs
        _set_run_text(r, t[:lo] + (new if not placed else "") + t[hi:])
        placed = True


def replace_paragraph_text(p_el: Any, new_text: str) -> bool:
    """Đặt text đoạn = new_text bằng diff tiền tố/hậu tố -> chỉ đụng vùng thay đổi (nhãn in đậm 'Bên A:' ... được giữ nguyên)."""
    old = _para_text(p_el)
    if old == new_text:
        return False
    if not _text_runs(p_el):
        return False
    p = 0
    while p < len(old) and p < len(new_text) and old[p] == new_text[p]:
        p += 1
    s = 0
    while s < len(old) - p and s < len(new_text) - p and old[len(old) - 1 - s] == new_text[len(new_text) - 1 - s]:
        s += 1
    replace_range(p_el, p, len(old) - s, new_text[p:len(new_text) - s])
    return True


_PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z0-9_\.]+)\s*\}\}")


def fill_placeholders(doc: DocumentType, values: dict[str, str]) -> tuple[list[dict], list[str]]:
    """Thay {{khoa}} bằng giá trị CRM (từ phải sang trái để chỉ số không lệch). Trả (đã thay, placeholder thiếu dữ liệu - để nguyên)."""
    applied: list[dict] = []
    missing: list[str] = []
    for i, el in enumerate(_all_paragraph_elements(doc)):
        text = _para_text(el)
        for m in reversed(list(_PLACEHOLDER.finditer(text))):
            key = m.group(1)
            val = values.get(key)
            if val is None or val == "":
                if key not in missing:
                    missing.append(key)
                continue
            replace_range(el, m.start(), m.end(), val)
            applied.append({"id": f"p{i}", "key": key, "value": val})
    return applied, missing


# ───────────────────────── bảng "nhãn: giá trị" của Bên A/Bên B (mẫu KHÔNG dùng {{}}) ─────────────────────────
# Bug thật: mẫu DOCX tái sử dụng từ 1 hợp đồng CŨ (của khách hàng khác) thường không có {{placeholder}} - MST/Địa chỉ/Người đại
# diện/Điện thoại/Email của Bên A/B nằm trong BẢNG dạng "nhãn | : giá trị" với chữ THẬT của khách hàng CŨ. fill_placeholders()
# không đụng tới bảng (không có {{}} để thay), và bước AI chỉnh đoạn CHỦ ĐỘNG bỏ qua mọi đoạn nằm trong bảng - nên dữ liệu
# khách hàng CŨ bị lọt nguyên vào hợp đồng MỚI (lộ MST/SĐT/email của khách khác). Hàm này tìm đúng các bảng 2 cột kiểu đó và
# ghi đè bằng dữ liệu Bên A/B THẬT của hợp đồng đang tạo.
def _fold(text: str) -> str:
    import unicodedata

    t = unicodedata.normalize("NFD", (text or "").lower().replace("đ", "d"))
    return "".join(c for c in t if unicodedata.category(c) != "Mn")


_SIDE_A_MARKERS = ("ben a", "ben yeu cau", "ben su dung dich vu", "ben mua")
_SIDE_B_MARKERS = ("ben b", "ben cung cap dich vu", "ben ban")
_PARTY_LABEL_KEYS = (
    ("tax_code", ("ma so thue", "mst")),
    ("address", ("dia chi",)),
    ("phone", ("dien thoai", "so dien thoai", "sdt")),
    ("email", ("email",)),
    ("position", ("chuc vu",)),
    ("rep", ("nguoi dai dien", "dai dien")),
)


def _party_label_key(label: str) -> str | None:
    f = _fold(label)
    for key, words in _PARTY_LABEL_KEYS:
        if any(f == w or f.startswith(w) for w in words):
            return key
    return None


_PARTY_BLANK = "………………………"


def _party_label_rows(tbl_el: Any) -> list[tuple[Any, str]]:
    """Các dòng [nhãn | giá trị] THỰC SỰ của bảng thông tin Bên A/B: cột nhãn khớp từ khoá đã biết VÀ cột giá trị KHÔNG
    tự nó cũng giống 1 nhãn (chặn bảng chữ ký kiểu 'ĐẠI DIỆN BÊN A (Ký, ghi rõ họ tên) | ĐẠI DIỆN BÊN B (...)' - 2 ô đều
    bắt đầu bằng 'đại diện' nên lọt qua match nhãn 'rep' nếu không có điều kiện này - bug đã gặp với file thật)."""
    rows = []
    for tr in tbl_el.findall(qn("w:tr")):
        tcs = tr.findall(qn("w:tc"))
        if len(tcs) < 2:
            continue
        key = _party_label_key(_cell_text(tcs[0]))
        if key and _party_label_key(_cell_text(tcs[1])) is None:
            rows.append((tr, key))
    return rows


def _fill_one_party_table(tbl_el: Any, party: dict[str, str] | None) -> list[dict]:
    if not party:
        return []
    label_rows = _party_label_rows(tbl_el)
    if len(label_rows) < 2:                 # bảng chữ ký/khác chỉ lọt 0-1 dòng "giống nhãn" - không phải bảng Bên A/B thật
        return []
    touched: list[dict] = []
    for tr, key in label_rows:
        tcs = tr.findall(qn("w:tc"))
        # Dữ liệu hợp đồng ĐANG tạo không có trường này (vd khách cá nhân không có MST) - để TRỐNG (………), KHÔNG được giữ
        # nguyên chữ của khách hàng CŨ còn sót lại trong mẫu tái sử dụng (chính là bug lộ thông tin khách khác đã báo).
        new_value = str(party.get(key) or "").strip() or _PARTY_BLANK
        value_paras = list(tcs[1].iter(qn("w:p")))
        if not value_paras:
            continue
        vp = value_paras[0]
        old_text = _para_text(vp)
        lead = (re.match(r"^\s*:?\s*", old_text) or re.match(r"", old_text)).group(0)
        if key == "rep":
            # O gop ca ten dai dien va "Chuc vu: ..." trong CUNG 1 o (vd "Ông Phan Văn Vũ   Chức vụ: Chủ hộ kinh doanh"), và tên
            # trong mau thuong nam o 1 run RIENG duoc in dam (de nhan manh), khac run cua nhan "Chuc vu:" phia sau. Dung
            # replace_paragraph_text() (diff tien to/hau to tren CA doan) o day se khien vung thay the tran qua ca 2 run, lam
            # nhan "Chuc vu:" bi "an theo" dinh dang in DAM cua run ten cu - bug da gap. Phai sua TUNG PHAN (ten / chuc vu)
            # bang replace_range() theo dung vi tri cu, xu ly PHAN SAU TRUOC de khong lam lech offset cua phan truoc.
            m = re.search(r"chức vụ\s*:", old_text, re.I)
            position = str(party.get("position") or "").strip()
            name_end = m.start() if m else len(old_text)
            changed = False
            if m:
                new_position_text = f"  Chức vụ: {position}" if position else f"  Chức vụ: {_PARTY_BLANK}"
                old_position_part = old_text[m.start():]
                if old_position_part != new_position_text:
                    replace_range(vp, m.start(), len(old_text), new_position_text)
                    changed = True
            elif position:
                replace_range(vp, len(old_text), len(old_text), f"   Chức vụ: {position}")
                changed = True
            old_name_part = old_text[len(lead):name_end].rstrip()
            if old_name_part != new_value:
                replace_range(vp, len(lead), name_end, new_value)
                changed = True
            if changed:
                # Phong truong hop run "Chuc vu: ..." bi gop chung run voi ten (vd mau chi co 1 run DUY NHAT in dam
                # cho ca dong, hoac replace_range() noi lien 2 doan vao cung run) - ep rieng phan nhan KHONG DAM,
                # du run goc dang dinh dang gi (bug "Chuc vu bi in dam" van gap du da sua truong hop run tach san).
                _unbold_label_suffix(vp, "Chức vụ")
                touched.append({"key": key, "old": old_text, "new": _para_text(vp)})
            continue
        new_text = f"{lead}{new_value}"
        if new_text != old_text and replace_paragraph_text(vp, new_text):
            touched.append({"key": key, "old": old_text, "new": new_text})
    return touched


def _fill_signature_names(tbl_el: Any, parties: dict[str, dict[str, str]]) -> list[dict]:
    """Dòng tên người ký ngay dưới '(Ký, ghi rõ họ tên...)' trong bảng chữ ký (vd 'ĐẠI DIỆN BÊN A ... PHAN VĂN VŨ') vẫn là
    chữ THẬT của khách hàng CŨ còn sót lại trong mẫu tái sử dụng - không nằm trong bảng nhãn:giá trị nên
    _fill_one_party_table() không đụng tới. Xác định Bên theo marker 'BÊN A/B' NGAY TRONG CÙNG Ô (bảng chữ ký thường
    2 cột A|B trên CÙNG 1 hàng, không thể dùng current_side theo thứ tự toàn tài liệu). Dòng tên trống (chưa ký) thì
    để nguyên, không tự điền."""
    touched: list[dict] = []
    for tc in tbl_el.iter(qn("w:tc")):
        paras = list(tc.iter(qn("w:p")))
        side: str | None = None
        name_idx: int | None = None
        for i, p in enumerate(paras):
            f = _fold(_para_text(p))
            if any(m in f for m in _SIDE_A_MARKERS):
                side = "a"
            elif any(m in f for m in _SIDE_B_MARKERS):
                side = "b"
            if "ky" in f and "ghi ro ho ten" in f:
                name_idx = i + 1
        if side is None or name_idx is None or name_idx >= len(paras):
            continue
        party = parties.get(side)
        if not party:
            continue
        name_p = paras[name_idx]
        old = _para_text(name_p)
        new_name = str(party.get("rep") or party.get("name") or "").strip()
        if new_name and old.strip() and old.strip() != new_name and replace_paragraph_text(name_p, new_name):
            touched.append({"key": "signature_name", "old": old, "new": new_name, "side": side})
    return touched


def fill_party_info_tables(doc: DocumentType, parties: dict[str, dict[str, str]] | None) -> dict[str, Any]:
    """Quét theo thứ tự xuất hiện: đoạn văn chứa 'BÊN A'/'BÊN B' (hoặc tương đương) xác định PHE đang nói tới, bảng 2 cột
    kiểu 'nhãn: giá trị' ngay sau đó được ghi đè bằng dữ liệu Bên tương ứng (parties['a']/['b'] từ build_parties() - cùng
    nguồn dữ liệu DOCX/PDF/Legal Check đang dùng, luôn nhất quán). Trả {'filled': [...], 'unresolvedTables': n} - n = số
    bảng giống bảng thông tin 2 bên nhưng không xác định được Bên nào (Sale cần tự kiểm tra)."""
    if not parties:
        return {"filled": [], "unresolvedTables": 0}
    filled: list[dict] = []
    unresolved = 0
    current_side: str | None = None
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            f = _fold(_para_text(child))
            if any(m in f for m in _SIDE_A_MARKERS):
                current_side = "a"
            elif any(m in f for m in _SIDE_B_MARKERS):
                current_side = "b"
        elif child.tag == qn("w:tbl"):
            party = parties.get(current_side) if current_side else None
            touched = _fill_one_party_table(child, party)
            if touched:
                filled.extend({**t, "side": current_side} for t in touched)
            elif current_side is None and len(_party_label_rows(child)) >= 2:
                unresolved += 1
            filled.extend(_fill_signature_names(child, parties))
    return {"filled": filled, "unresolvedTables": unresolved}


def apply_edits(doc: DocumentType, edits: Iterable[dict]) -> list[dict]:
    """Áp các đề xuất {id, text}. Trả danh sách thực sự đã đổi (kèm before/after)."""
    els = _all_paragraph_elements(doc)
    done: list[dict] = []
    for e in edits:
        m = re.fullmatch(r"p(\d+)", str(e.get("id") or ""))
        if not m or int(m.group(1)) >= len(els):
            continue
        el = els[int(m.group(1))]
        before = _para_text(el)
        if replace_paragraph_text(el, str(e.get("text") or "")):
            done.append({"id": e["id"], "before": before, "after": _para_text(el), "reason": e.get("reason", "")})
    return done


def replace_body_paragraphs(doc: DocumentType, new_texts: list[str]) -> list[dict]:
    """Thay TOÀN BỘ đoạn văn THƯỜNG (không nằm trong bảng, có chữ) bằng `new_texts` theo đúng thứ tự - CHO PHÉP
    thêm/bớt số đoạn tuỳ ý (khác apply_edits() chỉ thay text của đoạn ĐÃ CÓ, không tạo/xoá đoạn được). Dùng cho luồng
    "sửa tự do" 1 hợp đồng ĐÃ TẠO (không giới hạn số dòng) - bảng (hạng mục, thông tin Bên A/B...) không bị đụng tới,
    cùng khuôn với cách fill_items_table() nhân bản/xoá HÀNG bảng, áp dụng cho ĐOẠN VĂN thân tài liệu."""
    old_els = [el for el in _all_paragraph_elements(doc) if _para_text(el).strip() and not _in_table(el)]
    changed: list[dict] = []
    n = min(len(old_els), len(new_texts))
    for i in range(n):
        before = _para_text(old_els[i])
        if before != new_texts[i] and replace_paragraph_text(old_els[i], new_texts[i]):
            changed.append({"op": "update", "before": before, "after": new_texts[i]})
    if len(new_texts) > len(old_els):
        anchor = old_els[-1] if old_els else None
        for extra_text in new_texts[len(old_els):]:
            if anchor is None:                      # không có đoạn mẫu nào để nhân bản định dạng - bỏ qua phần dư
                break
            new_el = copy.deepcopy(anchor)
            replace_paragraph_text(new_el, extra_text)
            anchor.addnext(new_el)
            anchor = new_el
            changed.append({"op": "insert", "before": "", "after": extra_text})
    elif len(new_texts) < len(old_els):
        for el in old_els[len(new_texts):]:
            before = _para_text(el)
            el.getparent().remove(el)
            changed.append({"op": "delete", "before": before, "after": ""})
    return changed


_ARTICLE_HEADING_RE = re.compile(r"^\s*ĐIỀU\s+(\d+)\b(.*)$", re.I)


def renumber_articles(doc: DocumentType) -> int:
    """Đánh số lại TUẦN TỰ các dòng tiêu đề 'ĐIỀU n.' theo đúng thứ tự xuất hiện trong thân tài liệu (không tính đoạn
    trong bảng), giữ nguyên phần text SAU số. Gọi sau khi chèn/xoá điều khoản để số thứ tự luôn đúng, không trùng/lệch
    (feedback "Cho chèn Điều 8, 9, 10... vào đúng vị trí, tự cập nhật số thứ tự điều khoản khi cần"). Trả về số dòng đã đổi."""
    idx = 0
    changed = 0
    for el in _all_paragraph_elements(doc):
        if _in_table(el):
            continue
        text = _para_text(el)
        if not _ARTICLE_HEADING_RE.match(text.strip()):
            continue
        idx += 1
        new_text = re.sub(r"^(\s*)ĐIỀU\s+\d+", rf"\g<1>ĐIỀU {idx}", text, count=1, flags=re.I)
        if new_text != text and replace_paragraph_text(el, new_text):
            changed += 1
    return changed


def apply_structural_changes(doc: DocumentType, inserts: list[dict], deletes: list[dict]) -> dict:
    """Chèn ĐIỀU KHOẢN MỚI / xoá điều khoản theo id (đã qua validate_template_inserts/validate_template_deletes) -
    pipeline mẫu DOCX trước đây CHỈ sửa được nội dung đoạn có sẵn (apply_edits), không chèn/xoá được đoạn, khiến yêu
    cầu "thêm Điều 9" của Sale bị âm thầm bỏ qua (bug "Bước 2 không ăn yêu cầu").

    QUAN TRỌNG VỀ THỨ TỰ: id là CHỈ SỐ VỊ TRÍ (xem describe_paragraphs/paragraph_by_id) - resolve TOÀN BỘ id sang
    element TRƯỚC khi mutate bất kỳ gì (deepcopy() rồi addnext()/remove()), vì xoá/chèn làm LỆCH chỉ số các đoạn phía
    sau nếu resolve id sau khi đã mutate trước đó. Thứ tự an toàn: resolve hết -> XOÁ -> CHÈN (bằng element đã resolve,
    không resolve lại bằng id) -> đánh số lại 'ĐIỀU n.' nếu có thay đổi cấu trúc.

    `inserts`: [{"after_id": "p12"|"start"|None, "title": "ĐIỀU 9. ...", "body": "dòng 1\\ndòng 2"}, ...]
    `deletes`: [{"ids": ["p30", "p31"]}, ...]
    Trả {"inserted": [text...], "deletedCount": N, "renumbered": N} để lớp gọi báo lại cho Sale (không âm thầm)."""
    els = _all_paragraph_elements(doc)

    def el_by_id(pid: str | None):
        m = re.fullmatch(r"p(\d+)", str(pid or ""))
        return els[int(m.group(1))] if m and int(m.group(1)) < len(els) else None

    delete_targets: list[Any] = []
    for d in deletes or []:
        for pid in d.get("ids") or []:
            el = el_by_id(pid)
            if el is not None and el not in delete_targets:
                delete_targets.append(el)

    insert_specs: list[tuple[Any, str, dict]] = []
    needs_renumber = False
    for ins in inserts or []:
        after_id = ins.get("after_id")
        anchor_el = el_by_id(after_id) if after_id not in (None, "", "start") else None
        insert_mode = "after" if anchor_el is not None else "before"
        if anchor_el is None:
            anchor_el = next((e for e in els if not _in_table(e)), None)
            needs_renumber = True   # chen vao DAU body - truoc Dieu 1 hien co, chac chan lech so neu khong danh lai
        if anchor_el is not None and anchor_el not in delete_targets:
            insert_specs.append((anchor_el, insert_mode, ins))
            # Chen NGAY SAU dieu khoan cuoi cung (khong co "DIEU n." nao con lai phia sau) -> giu NGUYEN so Sale/AI da
            # ghi (vd "them Dieu 9" sau Dieu 7 hien co -> van la "Dieu 9", khong ep ve "Dieu 8"). Chen VAO GIUA (con
            # dieu khoan cu nam sau vi tri chen) -> PHAI danh so lai de tranh trung/lech so voi cac dieu phia sau.
            if not needs_renumber:
                seen_anchor = False
                for e in els:
                    if e is anchor_el:
                        seen_anchor = True
                        continue
                    if seen_anchor and not _in_table(e) and _ARTICLE_HEADING_RE.match(_para_text(e).strip()):
                        needs_renumber = True
                        break

    deleted_count = 0
    for el in delete_targets:
        if el.getparent() is not None:
            el.getparent().remove(el)
            deleted_count += 1

    inserted_texts: list[str] = []
    for anchor_el, insert_mode, ins in insert_specs:
        if anchor_el.getparent() is None:
            continue
        heading_el = None
        for el in els:
            if el is anchor_el:
                break
            if _in_table(el) or el.getparent() is None:
                continue
            if _ARTICLE_HEADING_RE.match(_para_text(el).strip()):
                heading_el = el
        if heading_el is None or heading_el.getparent() is None:
            heading_el = anchor_el
        title = str(ins.get("title") or "").strip()
        body_lines = [b for b in str(ins.get("body") or "").split("\n") if b.strip()]
        lines = ([title] if title else []) + body_lines
        cursor = anchor_el
        for i, text in enumerate(lines):
            template = heading_el if (i == 0 and title) else anchor_el
            new_el = copy.deepcopy(template)
            replace_paragraph_text(new_el, text)
            if insert_mode == "before" and i == 0:
                anchor_el.addprevious(new_el)
            else:
                cursor.addnext(new_el)
            cursor = new_el
            inserted_texts.append(text)

    if deleted_count:
        needs_renumber = True   # xoa doan co the de lai "lo hong" so thu tu, luon kiem tra lai
    renumbered = renumber_articles(doc) if needs_renumber else 0
    return {"inserted": inserted_texts, "deletedCount": deleted_count, "renumbered": renumbered}


# ───────────────────────── bảng hạng mục báo giá ─────────────────────────
# Thứ tự khoá QUAN TRỌNG: 'amount' (Thành tiền...) khớp trước 'vat' vì tiêu đề "Thành tiền (gồm VAT)" có chữ VAT; 'features' khớp trước 'desc' vì "Tính năng/mô tả".
_COL_KEYS = {
    "stt": ("stt", "tt", "no."),
    "amount": ("thành tiền", "thanh tien", "amount", "total", "phí dịch vụ", "phi dich vu", "giá trị", "gia tri"),
    "unit_price": ("đơn giá", "don gia", "unit price", "price"),
    "qty": ("số lượng", "so luong", "sl", "quantity", "qty"),
    "unit": ("đvt", "dvt", "đơn vị tính", "đơn vị", "unit"),
    "vat": ("vat", "thuế", "thue", "gtgt"),
    "features": ("tính năng", "tinh nang", "mô tả", "mo ta", "diễn giải", "dien giai", "quy cách", "thông số", "features", "specification"),
    "desc": ("hạng mục", "hang muc", "sản phẩm", "san pham", "dịch vụ", "dich vu", "tên", "nội dung", "noi dung", "description", "item"),
}


def item_total_incl_vat(item: dict) -> float:
    """Thành tiền (gồm VAT) của 1 hạng mục: totalAmount của báo giá; chỉ khi báo giá không có thì amountAfterDiscount x (1 + VAT%)."""
    ta, base = item.get("totalAmount"), float(item.get("amountAfterDiscount") or 0)
    if ta is not None and (float(ta) > 0 or base == 0):
        return float(ta)
    return base * (1 + float(item.get("vatRate") or 0) / 100)


_SUBTOTAL_LABEL = re.compile(r"(trước thuế|chưa (?:gồm |bao gồm )?(?:vat|thuế)|cộng tiền hàng|tổng tiền hàng|subtotal)", re.I)
_VAT_LABEL = re.compile(r"^\s*(thuế|tiền thuế|vat|thuế gtgt|gtgt)\b", re.I)
_TOTAL_LABEL = re.compile(r"^\s*(tổng cộng|tổng giá trị|tổng thanh toán|tổng tiền|tổng|cộng|total)\b", re.I)


def totals_kind(label: str) -> str | None:
    """'subtotal' | 'vat' | 'total' | None cho nhãn hàng tổng của bảng hạng mục."""
    if _SUBTOTAL_LABEL.search(label or ""):
        return "subtotal"
    if _VAT_LABEL.match(label or ""):
        return "vat"
    if _TOTAL_LABEL.match(label or ""):
        return "total"
    return None


def _cell_text(tc: Any) -> str:
    return " ".join(_para_text(p) for p in tc.iter(qn("w:p"))).strip()


def _classify_col(header: str) -> str | None:
    h = header.lower().strip()
    for key, words in _COL_KEYS.items():
        if any(h == w or (len(w) > 3 and w in h) for w in words):
            return key
    return None


def _table_columns(table) -> tuple[int, dict[str, int]] | None:
    for ri, row in enumerate(table.rows[:3]):
        cols: dict[str, int] = {}
        for ci, cell in enumerate(row.cells):
            k = _classify_col(cell.text)
            if k and k not in cols:
                cols[k] = ci
        if len(cols) >= 3 and "desc" in cols and ("amount" in cols or "unit_price" in cols):
            return ri, cols
    return None


def list_item_tables(doc: DocumentType) -> list[dict[str, Any]]:
    """Các bảng của mẫu có tiêu đề nhận diện được là bảng hạng mục (để người dùng chọn khi mẫu có nhiều bảng)."""
    out = []
    for idx, table in enumerate(doc.tables):
        found = _table_columns(table)
        if found:
            out.append({"index": idx, "rows": len(table.rows), "columns": sorted(found[1], key=found[1].get),
                        "preview": " | ".join(c.text.strip() for c in table.rows[found[0]].cells)[:120]})
    return out


def find_items_table(doc: DocumentType, table_index: int | None = None) -> tuple[Any, int, dict[str, int]] | None:
    """Bảng hạng mục: theo chỉ số người dùng chọn, hoặc bảng nhận diện được đầu tiên. KHÔNG tự chèn bảng mới."""
    for idx, table in enumerate(doc.tables):
        if table_index is not None and idx != table_index:
            continue
        found = _table_columns(table)
        if found:
            return table, found[0], found[1]
    return None


def _format_number(value: Any, currency: str = "VND") -> str:
    try:
        n = float(value)
    except (TypeError, ValueError):
        return ""
    if (currency or "VND").upper() == "VND":
        return f"{int(round(n)):,}".replace(",", ".")
    return f"{n:,.2f}"


def _set_cell(tc: Any, text: str) -> None:
    ps = list(tc.iter(qn("w:p")))
    if not ps:
        return
    replace_paragraph_text(ps[0], text)
    for extra in ps[1:]:                          # ô nhiều đoạn: chỉ giữ đoạn đầu mang giá trị
        for t in extra.iter(qn("w:t")):
            t.text = ""


def _tc_at_logical_col(tr: Any, col_index: int) -> Any | None:
    """Ô <w:tc> THỰC SỰ nằm ở cột logic `col_index` (0-based, tính theo header) - hàng tổng thường GỘP Ô nhãn (vd
    'TỔNG CỘNG' chiếm 2 cột đầu qua gridSpan=2), nên <w:tc> CUỐI CÙNG trong XML KHÔNG chắc là cột cuối cùng của
    header (bug đã gặp: hàng tổng 4 cột STT/Hạng mục/Phí dịch vụ/Đơn vị, ô nhãn gộp STT+Hạng mục -> chỉ còn 3 <w:tc>,
    cells[-1] lại trúng ô 'Đơn vị' thay vì ô 'Phí dịch vụ' - tiền tổng ghi nhầm sang cột đơn vị, cột tiền giữ số cũ)."""
    pos = 0
    for tc in tr.findall(qn("w:tc")):
        tcpr = tc.find(qn("w:tcPr"))
        span_el = tcpr.find(qn("w:gridSpan")) if tcpr is not None else None
        span = int(span_el.get(qn("w:val"))) if span_el is not None else 1
        if pos <= col_index < pos + span:
            return tc
        pos += span
    return None


def fill_items_table(doc: DocumentType, quote: dict | None, table_index: int | None = None) -> dict[str, Any]:
    """Điền TOÀN BỘ hạng mục báo giá (đúng thứ tự, không cắt) vào bảng của mẫu: nhân bản HÀNG MẪU đầu tiên (giữ border/shading/font).
    Giá trị 100% từ báo giá CRM. Hàng tổng: 'trước thuế' / 'VAT' / 'tổng thanh toán' lấy từ báo giá; hàng không nhận diện giữ nguyên + cảnh báo.
    Không tìm được bảng -> KHÔNG chèn gì, trả `candidates` để người dùng chọn."""
    result: dict[str, Any] = {"filled": False, "rows": 0, "warnings": [], "candidates": list_item_tables(doc), "tableIndex": None}
    items = list((quote or {}).get("items") or [])
    found = find_items_table(doc, table_index)
    if not found:
        if items:
            result["warnings"].append("Mẫu không có bảng hạng mục nhận diện được - hạng mục báo giá KHÔNG được chèn tự động. Hãy chọn bảng cần điền (nếu có) hoặc bổ sung bảng vào mẫu.")
        return result
    if not items:
        result["warnings"].append("Chưa có báo giá/hạng mục - bảng hạng mục giữ nguyên như mẫu.")
        return result
    table, header_idx, cols = found
    result["tableIndex"] = next(i for i, t in enumerate(doc.tables) if t._tbl is table._tbl)
    currency = (quote.get("currency") or "VND").upper()
    tr_list = table._tbl.findall(qn("w:tr"))
    data_rows = tr_list[header_idx + 1:]
    if not data_rows:
        result["warnings"].append("Bảng hạng mục không có hàng mẫu để nhân bản.")
        return result

    def row_label(tr: Any) -> str:
        return next((_cell_text(c) for c in tr.findall(qn("w:tc")) if _cell_text(c)), "")

    body_rows = [r for r in data_rows if totals_kind(row_label(r)) is None]
    total_rows = [r for r in data_rows if totals_kind(row_label(r)) is not None]
    proto = copy.deepcopy(body_rows[0] if body_rows else data_rows[0])
    anchor = body_rows[0] if body_rows else data_rows[0]
    for r in body_rows:
        if r is not anchor:
            r.getparent().remove(r)
    n = 0
    for item in items:
        row = copy.deepcopy(proto)
        tcs = row.findall(qn("w:tc"))
        desc = str(item.get("description") or item.get("serviceDescription") or "").strip()
        if item.get("rowType") == "section":                       # tiêu đề nhóm: chỉ ô tên hạng mục, các ô còn lại để trống
            values = {k: "" for k in cols}
            values["desc"] = desc
        else:
            n += 1
            feat = str(item.get("serviceDescription") or "").strip()
            vat = item.get("vatRate")
            values = {
                "stt": str(n), "desc": desc, "features": "" if feat == desc else feat, "unit": str(item.get("unit") or ""),
                "qty": (f"{int(item['quantity'])}" if float(item.get("quantity") or 0).is_integer() else str(item.get("quantity")).replace(".", ",")) if item.get("quantity") is not None else "",
                "unit_price": _format_number(item.get("unitPrice"), currency) if item.get("unitPrice") is not None else "",
                "vat": f"{vat:g}%" if vat is not None else "", "amount": _format_number(item_total_incl_vat(item), currency),
            }
        for key, ci in cols.items():
            if ci < len(tcs) and key in values:
                _set_cell(tcs[ci], values[key])
        anchor.addprevious(row)
    anchor.getparent().remove(anchor)
    sums = {"subtotal": (quote or {}).get("subtotalAmount"), "vat": (quote or {}).get("vatAmount"), "total": (quote or {}).get("totalAmount")}
    for tr in total_rows:
        label = row_label(tr)
        kind = totals_kind(label)
        amount_tc = _tc_at_logical_col(tr, cols["amount"]) if "amount" in cols else None
        if kind and sums.get(kind) is not None and amount_tc is not None:
            _set_cell(amount_tc, _format_number(sums[kind], currency))
        else:
            result["warnings"].append(f"Hàng '{label[:40]}' trong bảng hạng mục giữ nguyên theo mẫu - cần đối chiếu thủ công.")
    result.update({"filled": True, "rows": n})
    return result


def docx_to_clauses(docx_bytes: bytes) -> list[dict]:
    """Nội dung THỰC của DOCX -> [{title, body}] tách theo 'ĐIỀU n' (gồm cả chữ trong bảng theo đúng thứ tự). Dùng để AI kiểm tra rủi ro đúng tài liệu đã lưu."""
    doc = load_document(docx_bytes)
    body = doc.element.body
    clauses: list[dict] = []
    cur = {"title": "Phần mở đầu", "lines": []}
    head = re.compile(r"^\s*ĐIỀU\s+\d+", re.I)
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            text = _para_text(child).strip()
            if not text:
                continue
            if head.match(text):
                clauses.append({"title": cur["title"], "body": "\n".join(cur["lines"])})
                cur = {"title": text, "lines": []}
            else:
                cur["lines"].append(text)
        elif child.tag == qn("w:tbl"):
            for tr in child.iter(qn("w:tr")):
                cells = [_cell_text(tc) for tc in tr.findall(qn("w:tc"))]
                # BUG THAT DA GAP: loc bo o trong roi moi "|".join lam LECH COT khi 1 o giua bang trong (vd cot
                # "Tinh nang/mo ta" trung voi "Hang muc" nen bi rong) - AI doc nham gia tri cot sau thanh cot truoc
                # (vd don gia bi hieu thanh SL). Giu DUNG vi tri cot bang dau "-" cho o trong, chi bo qua dong trong HOAN TOAN.
                if any(c.strip() for c in cells):
                    line = " | ".join(c.strip() or "—" for c in cells)
                    cur["lines"].append(line)
    clauses.append({"title": cur["title"], "body": "\n".join(cur["lines"])})
    return [c for c in clauses if c["body"].strip() or c["title"] != "Phần mở đầu"]


# ───────────────────────── xuất bytes ─────────────────────────

def save_document(doc: DocumentType) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ───────────────────────── DOCX -> PDF (LibreOffice) ─────────────────────────

def find_soffice() -> str | None:
    explicit = os.environ.get("CONTRACT_SOFFICE_PATH")
    if explicit and os.path.exists(explicit):
        return explicit
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    return None


# ── giới hạn tài nguyên / an toàn khi chạy LibreOffice ──
def _env_int(name: str, default: int) -> int:
    try:
        return max(1, int(os.environ.get(name, default)))
    except ValueError:
        return default


MAX_DOCX_BYTES = 10 * 1024 * 1024            # file nạp vào
MAX_DOCX_UNCOMPRESSED = 120 * 1024 * 1024    # chặn zip-bomb
MAX_DOCX_ENTRIES = 5000
MAX_PDF_BYTES = 100 * 1024 * 1024            # PDF sinh ra (RLIMIT_FSIZE + kiểm tra lại)
_slots = threading.BoundedSemaphore(_env_int("CONTRACT_PDF_MAX_CONCURRENCY", 2))   # mỗi tiến trình soffice ~300-500MB RAM


def validate_docx_bytes(content: bytes) -> None:
    """Từ chối file quá lớn / không phải DOCX / zip-bomb TRƯỚC khi đưa vào python-docx hay LibreOffice."""
    if len(content) > MAX_DOCX_BYTES:
        raise ValueError("File DOCX quá lớn (tối đa 10MB).")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            infos = z.infolist()
            if len(infos) > MAX_DOCX_ENTRIES or sum(i.file_size for i in infos) > MAX_DOCX_UNCOMPRESSED:
                raise ValueError("File DOCX có cấu trúc bất thường (quá nhiều/quá lớn thành phần).")
            if "word/document.xml" not in z.namelist():
                raise ValueError("File không phải tài liệu Word (.docx) hợp lệ.")
    except zipfile.BadZipFile as exc:
        raise ValueError("File không phải tài liệu Word (.docx) hợp lệ.") from exc


def _kill_tree(proc: subprocess.Popen) -> None:
    """soffice là launcher sinh ra soffice.bin: phải diệt CẢ nhóm tiến trình, không để mồ côi sau timeout."""
    try:
        if os.name == "posix":
            os.killpg(proc.pid, signal.SIGKILL)
        else:
            proc.kill()
    except (ProcessLookupError, PermissionError, OSError):
        pass


def convert_docx_to_pdf(docx_bytes: bytes, timeout: int | None = None) -> bytes:
    soffice = find_soffice()
    if not soffice:
        raise PdfEngineUnavailable(
            "Server chưa cài LibreOffice (soffice) nên chưa xuất được PDF theo đúng mẫu DOCX. "
            "Hãy tải DOCX đã chỉnh sửa; PDF sẽ khả dụng khi backend chạy bằng Docker image có LibreOffice."
        )
    validate_docx_bytes(docx_bytes)
    timeout = timeout or _env_int("CONTRACT_SOFFICE_TIMEOUT", 90)
    if not _slots.acquire(timeout=timeout):
        raise PdfEngineUnavailable("Đang có quá nhiều yêu cầu xuất PDF cùng lúc, vui lòng thử lại sau ít phút.")
    try:
        with tempfile.TemporaryDirectory(prefix="contract_pdf_") as tmp:        # mkdtemp: quyền 0700, luôn bị xoá kể cả khi lỗi/timeout
            src = os.path.join(tmp, "contract.docx")
            with open(src, "wb") as f:
                f.write(docx_bytes)
            profile = os.path.join(tmp, "profile")                              # profile riêng mỗi lần: không tranh khoá, không rò dữ liệu giữa các yêu cầu
            cmd = [soffice, f"-env:UserInstallation=file:///{profile.replace(os.sep, '/').lstrip('/')}", "--headless", "--norestore", "--nolockcheck",
                   "--nodefault", "--nologo", "--convert-to", "pdf:writer_pdf_Export", "--outdir", tmp, src]
            env = {**os.environ, "HOME": tmp, "TMPDIR": tmp, "SAL_USE_VCLPLUGIN": "svp"}
            proc = subprocess.Popen(cmd, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env,
                                    start_new_session=(os.name == "posix"))
            try:
                if os.name == "posix":
                    import resource
                    resource.prlimit(proc.pid, resource.RLIMIT_CPU, (timeout * 2, timeout * 2))
                    resource.prlimit(proc.pid, resource.RLIMIT_FSIZE, (MAX_PDF_BYTES, MAX_PDF_BYTES))
            except Exception:  # noqa: BLE001  (không phải Linux/không đủ quyền: vẫn còn timeout + kill)
                pass
            try:
                _out, err = proc.communicate(timeout=timeout)
            except subprocess.TimeoutExpired as exc:
                _kill_tree(proc)
                proc.communicate()
                raise PdfEngineUnavailable("LibreOffice chuyển đổi PDF quá thời gian cho phép.") from exc
            finally:
                _kill_tree(proc)                                                # dọn soffice.bin còn sót
            out = os.path.join(tmp, "contract.pdf")
            if proc.returncode != 0 or not os.path.exists(out):
                raise PdfEngineUnavailable("LibreOffice chuyển đổi PDF thất bại: %r" % ((err or b"")[:300],))
            if os.path.getsize(out) > MAX_PDF_BYTES:
                raise PdfEngineUnavailable("PDF sinh ra quá lớn.")
            with open(out, "rb") as f:
                return f.read()
    finally:
        _slots.release()


_health_cache: dict[str, Any] = {"at": 0.0, "data": None}


def health_check(deep: bool = False, cache_seconds: int = 300) -> dict[str, Any]:
    """Xác nhận engine PDF dùng được. Nông: tìm soffice + `--version` + font tiếng Việt. Sâu: chuyển thử 1 DOCX tiếng Việt -> PDF và đọc lại text."""
    soffice = find_soffice()
    result: dict[str, Any] = {"ok": False, "soffice": soffice, "version": None, "fontsVietnamese": None, "deep": None}
    if not soffice:
        result["error"] = "Không tìm thấy soffice (LibreOffice) trong PATH."
        return result
    try:
        proc = subprocess.run([soffice, "--headless", "--version"], capture_output=True, timeout=30, stdin=subprocess.DEVNULL,
                              env={**os.environ, "HOME": tempfile.gettempdir()})
        result["version"] = proc.stdout.decode("utf8", "ignore").strip() or None
    except Exception as exc:  # noqa: BLE001
        result["error"] = "soffice --version lỗi: %s" % exc
        return result
    fams = {f.lower() for f in _installed_fonts()}
    result["fontsVietnamese"] = bool(fams & {"liberation serif", "dejavu serif"}) if fams else None
    result["ok"] = bool(result["version"]) and result["fontsVietnamese"] is not False
    if deep:
        now = time.time()
        if _health_cache["data"] and now - _health_cache["at"] < cache_seconds:
            result["deep"] = _health_cache["data"]
        else:
            try:
                t0 = time.time()
                docx = build_docx_from_clauses("Kiểm tra", "HC-1", [{"title": "ĐIỀU 1. TIẾNG VIỆT", "body": "Đặng Thị Hương — Phạm Quỳnh Như — Nguyễn Hữu Thọ."}])
                pdf = convert_docx_to_pdf(docx, timeout=60)
                text = pdf_text(pdf)
                ok = "ĐIỀU 1" in text and "Đặng Thị Hương" in text
                result["deep"] = {"ok": ok, "seconds": round(time.time() - t0, 2), "pdfBytes": len(pdf), "vietnameseTextOk": ok}
            except Exception as exc:  # noqa: BLE001
                result["deep"] = {"ok": False, "error": str(exc)[:200]}
            _health_cache.update(at=now, data=result["deep"])
        result["ok"] = result["ok"] and bool(result["deep"] and result["deep"].get("ok"))
    return result


# Font có tương thích số liệu (metric-compatible) -> không làm lệch bố cục khi LibreOffice thay thế
METRIC_COMPATIBLE = {
    "times new roman": "Liberation Serif", "arial": "Liberation Sans", "courier new": "Liberation Mono",
    "calibri": "Carlito", "cambria": "Caladea",
}


def fonts_used(docx_bytes: bytes) -> list[str]:
    """Font THỰC SỰ ảnh hưởng bố cục: font trong thân tài liệu, header/footer, docDefaults và style Normal (bỏ qua font của các style list/ít dùng)."""
    names: set[str] = set()

    def collect(xml: str) -> None:
        for tag in re.findall(r"<w:rFonts\b[^>]*>", xml):
            names.update(re.findall(r'w:(?:ascii|hAnsi|cs|eastAsia)="([^"]+)"', tag))

    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        for part in z.namelist():
            if re.fullmatch(r"word/(document|header\d*|footer\d*)\.xml", part):
                collect(z.read(part).decode("utf8", "ignore"))
            elif part == "word/styles.xml":
                xml = z.read(part).decode("utf8", "ignore")
                for block in re.findall(r"<w:docDefaults>.*?</w:docDefaults>", xml, flags=re.S):
                    collect(block)
                for block in re.findall(r'<w:style\b[^>]*w:styleId="Normal"[^>]*>.*?</w:style>', xml, flags=re.S):
                    collect(block)
    return sorted(n for n in names if n)


def font_report(docx_bytes: bytes, installed: Iterable[str] | None = None) -> dict[str, Any]:
    """Báo font dùng trong DOCX mà server KHÔNG có (LibreOffice sẽ thay thế -> có thể lệch bố cục). Không thay thế âm thầm."""
    if installed is None:
        installed = _installed_fonts()
        if not installed:                                   # không có fc-list (vd Windows dev): không kết luận
            return {"missing": [], "metricCompatible": [], "unknown": True}
    have = {f.lower() for f in installed}
    missing, via_metric = [], []
    for name in fonts_used(docx_bytes):
        low = name.lower()
        if low in have:
            continue
        sub = METRIC_COMPATIBLE.get(low)
        if sub and sub.lower() in have:
            via_metric.append(f"{name} -> {sub}")
        else:
            missing.append(name)
    return {"missing": missing, "metricCompatible": via_metric, "unknown": False}


def _installed_fonts() -> list[str]:
    try:
        out = subprocess.run(["fc-list", ":", "family"], capture_output=True, timeout=20).stdout.decode("utf8", "ignore")
    except Exception:  # noqa: BLE001
        return []
    fams: set[str] = set()
    for line in out.splitlines():
        for f in line.split(","):
            fams.add(f.strip())
    return sorted(fams)


# ───────────────────────── PDF thống kê / phân loại mẫu PDF ─────────────────────────

def pdf_page_count(pdf_bytes: bytes) -> int:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        return len(pdf.pages)


def pdf_text(pdf_bytes: bytes) -> str:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        return "\n".join((p.extract_text() or "") for p in pdf.pages)


def classify_pdf_template(pdf_bytes: bytes) -> dict[str, Any]:
    """Phân biệt PDF có text / scan / có form field (AcroForm). Engine KHÔNG chỉnh layout PDF -> luôn layoutPreserved=False."""
    import pdfplumber
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        pages = len(pdf.pages)
        chars = sum(len(p.chars) for p in pdf.pages)
        images = sum(len(p.images) for p in pdf.pages)
        has_form = False
        try:
            catalog = pdf.doc.catalog
            has_form = "AcroForm" in catalog
        except Exception:  # noqa: BLE001
            pass
    if has_form:
        kind = "form"
        note = ("PDF có form field (AcroForm). Hệ thống CHƯA hỗ trợ điền trường PDF tự động - "
                "hãy chuyển mẫu sang DOCX hoặc điền form PDF bằng công cụ chuyên dụng.")
    elif chars < 40 * max(pages, 1) and images:
        kind = "scan"
        note = "PDF là ảnh scan (không có lớp text) - không thể chỉnh sửa hay trích nội dung. Hãy cung cấp bản DOCX hoặc PDF dạng text."
    else:
        kind = "text"
        note = ("PDF dạng text: AI chỉ tham chiếu văn phong, KHÔNG giữ nguyên form PDF. "
                "Muốn giữ đúng bố cục mẫu hãy chuyển sang DOCX; bản xuất ra sẽ theo bố cục chuẩn của hệ thống.")
    return {"kind": kind, "pages": pages, "textChars": chars, "layoutPreserved": False, "warning": note}


# ───────────────────────── dựng DOCX chuẩn từ điều khoản (không có mẫu DOCX) ─────────────────────────

def build_docx_from_clauses(title: str, contract_number: str, clauses: list[dict], party_a: str = "", party_b: str = "") -> bytes:
    """Dùng KHI người dùng không chọn mẫu DOCX: dựng hợp đồng Times New Roman 13pt, A4, quốc hiệu, chữ ký 2 bên - DOCX là nguồn chuẩn
    để xuất PDF qua LibreOffice (thay cho PDF tự dựng không dấu/1 trang trước đây)."""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(2), Cm(2), Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(13)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Times New Roman")

    def para(text: str, bold: bool = False, align=None, space_after: int = 6, italic: bool = False):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        return p

    C, J = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY
    para("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, align=C, space_after=0)
    para("Độc lập - Tự do - Hạnh phúc", bold=True, align=C, space_after=12)
    para(title.upper(), bold=True, align=C, space_after=2)
    if contract_number:
        para(f"Số: {contract_number}", align=C, space_after=12)
    for c in clauses:
        para(str(c.get("title") or ""), bold=True, space_after=3)
        for line in str(c.get("body") or "").split("\n"):
            if line.strip():
                para(line.strip(), align=J, space_after=4)
    para("")
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ci, (who, name) in enumerate((("ĐẠI DIỆN BÊN A", party_a), ("ĐẠI DIỆN BÊN B", party_b))):
        top, bottom = table.cell(0, ci), table.cell(1, ci)
        top.text = who
        top.paragraphs[0].alignment = C
        top.paragraphs[0].runs[0].bold = True
        bottom.text = f"(Ký, ghi rõ họ tên, đóng dấu)\n\n\n{name}"
        bottom.paragraphs[0].alignment = C
    return save_document(doc)
