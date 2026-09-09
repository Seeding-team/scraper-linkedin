"""Import "Bảng giá VPS Zone" (SecurityZone) từ file nguồn
"260623_v1_CPC[2026]_Gói Claude(1).xlsx" — đọc trực tiếp 2 sheet nguồn
("1.IN-Software AI" + "EU" cho mục cha I, "2.IN-Software AI" cho mục cha II)
bằng openpyxl, KHÔNG hardcode số đã tính tay (chỉ hardcode STT/cột cần đọc).

An toàn theo mặc định: LUÔN chạy --dry-run (chỉ in ra sẽ tạo/bỏ qua gì, không
ghi DB) trừ khi truyền --apply. Idempotent nhờ UNIQUE(price_book_version_id,
source_sheet, source_stt) ở DB — chạy lại --apply không tạo trùng.

Yêu cầu trước khi --apply: migration 106_price_book_vps_zone.sql đã áp dụng
lên DB thật (xem plan). KHÔNG tự áp dụng migration từ script này.

Loại bỏ: dòng "Dịch vụ triển khai" (STT=0, mục cha rỗng "XVI", giá 0đ, category
"Triển khai" khác "Software AI") — không phải sản phẩm Claude thật.

Usage:
    python scripts/import_price_book_vps_zone.py --xlsx "C:\\path\\to\\file.xlsx"
    python scripts/import_price_book_vps_zone.py --xlsx "..." --apply
"""

from __future__ import annotations

import argparse
import sys
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402

import app.core.config  # noqa: E402,F401 - phai import truoc de load .env/.env.local (dotenv), neu khong get_supabase_client() se bao thieu SUPABASE_URL khi chay script doc lap (khong qua app.main).
from app.core.supabase_client import get_supabase_client  # noqa: E402

EXCHANGE_RATE = Decimal("26326")
ISSUER_CODE = "SZ"  # SecurityZone — đã tồn tại sẵn trong quote_issuer_companies
PRICE_BOOK_CODE = "VPS_ZONE"
PRICE_BOOK_NAME = "Bảng giá VPS Zone"

# Nhóm I: nguồn "1.IN-Software AI" (giá vốn USD) + "EU" (giá bán/Rate/mô tả/giá tham chiếu).
# STT khớp giữa 2 sheet. import_duty_percent = 0 cho cả 4 dòng (đã verify:
# 66.66 * 26326 = 1.754.891,16 khớp tuyệt đối golden fixture, KHÔNG nhân thêm
# thuế nhập khẩu dù tham số chung ở sheet TONG HOP ghi 10%).
GROUP_I_STT = ["1", "2", "3", "4"]
GROUP_I_SKU_MAP = {"1": "ZONE-I-01", "2": "ZONE-I-02", "3": "ZONE-I-03", "4": "ZONE-I-04"}

# Nhóm II: nguồn "2.IN-Software AI" (VND trực tiếp) + "EU" (STT 201-205, EU=IN, VAT=0%).
GROUP_II_STT = ["201", "202", "203", "204", "205"]
GROUP_II_SKU_MAP = {
    "201": "ZONE-II-201",
    "202": "ZONE-II-202",
    "203": "ZONE-II-203",
    "204": "ZONE-II-204",
    "205": "ZONE-II-205",
}

EXCLUDED_STT_NOTE = "STT=0 'Dịch vụ triển khai' (mục cha rỗng 'XVI', giá 0đ) — loại, không phải sản phẩm Claude thật."


def _num(value) -> Decimal | None:
    if value is None or value == "":
        return None
    return Decimal(str(value))


def _read_group_i(wb) -> list[dict]:
    ws_in = wb["1.IN-Software AI"]
    ws_eu = wb["EU"]

    in_rows: dict[str, dict] = {}
    for r in range(2, 6):  # STT 1-4 ở hàng 2-5
        stt = str(int(ws_in.cell(row=r, column=1).value))
        if stt not in GROUP_I_STT:
            continue
        in_rows[stt] = {
            "name": ws_in.cell(row=r, column=3).value,  # C: Description Items (tên ngắn, đúng nguồn per-item)
            "unit_price_usd": _num(ws_in.cell(row=r, column=6).value),  # F: Unit Price (USD)
            "vat_in_percent": (_num(ws_in.cell(row=r, column=11).value) or Decimal(0)) * 100,  # K: % VAT
        }

    eu_rows: dict[str, dict] = {}
    for r in range(3, 7):  # STT 1-4 ở hàng 3-6 trong sheet EU
        stt = str(int(ws_eu.cell(row=r, column=1).value))
        if stt not in GROUP_I_STT:
            continue
        amount_before_vat = _num(ws_eu.cell(row=r, column=6).value)  # F: Thành Tiền VND (truoc VAT)
        vat_amount = _num(ws_eu.cell(row=r, column=7).value)  # G: Thuế VAT (SO TIEN, khong phai %)
        vat_eu_percent = (vat_amount / amount_before_vat * 100) if amount_before_vat else Decimal(0)
        eu_rows[stt] = {
            "description": ws_eu.cell(row=r, column=12).value,  # L: Description Items (mô tả dài nguyên văn)
            "unit_price_vnd": _num(ws_eu.cell(row=r, column=5).value),  # E: Đơn giá VND (giá bán trước VAT)
            "vat_eu_percent": vat_eu_percent,  # tinh tu G/F (sheet khong co cot % rieng)
            "rate_percent": (_num(ws_eu.cell(row=r, column=10).value) or Decimal(0)) * 100,  # J: Rate
            "reference_price": _num(ws_eu.cell(row=r, column=15).value),  # O: Đơn giá tham chiếu
            "reference_link": ws_eu.cell(row=r, column=20).value,  # T: Link Dự án tham chiếu
        }

    items = []
    for stt in GROUP_I_STT:
        merged = {**in_rows[stt], **eu_rows[stt]}
        items.append(
            {
                "source_sheet": "1.IN-Software AI",
                "source_stt": stt,
                "sku": GROUP_I_SKU_MAP[stt],
                "name": merged["name"],
                "description": merged["description"],
                "unit": "Gói",
                "default_quantity": Decimal(1),
                "cost_mode": "usd",
                "unit_price_usd": merged["unit_price_usd"],
                "exchange_rate": EXCHANGE_RATE,
                "import_duty_percent": Decimal(0),
                "vat_in_percent": merged["vat_in_percent"],
                "default_rate_percent": merged["rate_percent"],
                "vat_eu_percent": merged["vat_eu_percent"],
                "reference_price": merged["reference_price"],
                "reference_link": merged["reference_link"],
                "unit_price_vnd_direct": None,
                "vendor_name": None,
                "quote_link": None,
                "list_price_usd": None,
                "quote_received_date": None,
            }
        )
    return items


def _read_group_ii(wb) -> list[dict]:
    ws_in = wb["2.IN-Software AI"]
    ws_eu = wb["EU"]

    in_rows: dict[str, dict] = {}
    for r in range(2, 7):  # STT 201-205 ở hàng 2-6
        stt = str(int(ws_in.cell(row=r, column=1).value))
        if stt not in GROUP_II_STT:
            continue
        in_rows[stt] = {
            "name": ws_in.cell(row=r, column=3).value,  # C: Description Items (tên ngắn, đúng nguồn per-item)
            "unit_price_vnd": _num(ws_in.cell(row=r, column=8).value),  # H: Đơn giá VND
        }

    eu_rows: dict[str, dict] = {}
    for r in range(8, 13):  # STT 201-205 ở hàng 8-12 trong sheet EU
        stt = str(int(ws_eu.cell(row=r, column=1).value))
        if stt not in GROUP_II_STT:
            continue
        eu_rows[stt] = {
            "description": ws_eu.cell(row=r, column=12).value or ws_eu.cell(row=r, column=3).value,  # L hoặc C
        }

    items = []
    for stt in GROUP_II_STT:
        merged = {**in_rows[stt], **eu_rows[stt]}
        items.append(
            {
                "source_sheet": "2.IN-Software AI",
                "source_stt": stt,
                "sku": GROUP_II_SKU_MAP[stt],
                "name": merged["name"],
                "description": merged["description"],
                "unit": "Gói",
                "default_quantity": Decimal(1),
                "cost_mode": "vnd",
                "unit_price_vnd_direct": merged["unit_price_vnd"],
                "import_duty_percent": Decimal(0),
                "vat_in_percent": Decimal(0),
                "default_rate_percent": Decimal(0),
                "vat_eu_percent": Decimal(0),
                "reference_price": None,
                "reference_link": None,
                "unit_price_usd": None,
                "exchange_rate": None,
                "vendor_name": None,
                "quote_link": None,
                "list_price_usd": None,
                "quote_received_date": None,
            }
        )
    return items


def _decimal_to_jsonable(value):
    if isinstance(value, Decimal):
        return float(value)
    return value


def _get_admin_actor_id(client) -> str | None:
    result = client.table("app_users").select("id").eq("email", "admin@gmail.com").limit(1).execute()
    return result.data[0]["id"] if result.data else None


def _get_issuer_company_id(client) -> str:
    result = client.table("quote_issuer_companies").select("id").eq("code", ISSUER_CODE).limit(1).execute()
    if not result.data:
        raise SystemExit(f"Khong tim thay quote_issuer_companies voi code={ISSUER_CODE!r} - kiem tra lai migration 069.")
    return result.data[0]["id"]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", required=True, help="Duong dan file 260623_v1_CPC[2026]_Goi Claude(1).xlsx")
    parser.add_argument("--apply", action="store_true", help="Ghi that vao DB (mac dinh chi dry-run)")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    group_i = _read_group_i(wb)
    group_ii = _read_group_ii(wb)
    all_items = group_i + group_ii

    print(f"Doc duoc {len(all_items)} san pham tu file (bo qua 'Dich vu trien khai': {EXCLUDED_STT_NOTE})")
    for item in all_items:
        print(f"  - {item['sku']:12s} | {item['source_sheet']:20s} STT={item['source_stt']:>4s} | {item['name']}")

    if not args.apply:
        print("\n[DRY-RUN] Khong ghi DB. Chay lai voi --apply de ghi that (yeu cau migration 106 da duoc ap dung).")
        return 0

    client = get_supabase_client()
    actor_id = _get_admin_actor_id(client)
    issuer_company_id = _get_issuer_company_id(client)

    price_book = client.table("price_books").select("*").eq("code", PRICE_BOOK_CODE).limit(1).execute()
    if price_book.data:
        price_book_id = price_book.data[0]["id"]
        print(f"Da co price_book {PRICE_BOOK_CODE} -> {price_book_id}")
    else:
        created = client.table("price_books").insert(
            {
                "issuer_company_id": issuer_company_id,
                "code": PRICE_BOOK_CODE,
                "name": PRICE_BOOK_NAME,
                "status": "active",
                "created_by": actor_id,
            }
        ).execute()
        price_book_id = created.data[0]["id"]
        print(f"Da tao price_book {PRICE_BOOK_CODE} -> {price_book_id}")

    draft_version = client.rpc(
        "price_book_create_or_reuse_draft", {"p_price_book_id": price_book_id, "p_actor_id": actor_id}
    ).execute()
    version_row = draft_version.data
    if isinstance(version_row, list):
        version_row = version_row[0]
    version_id = version_row["id"]
    print(f"Draft version dang dung: version={version_row['version']} id={version_id}")

    section_ids: dict[str, str] = {}
    for group_label, source_sheet in (("I", "1.IN-Software AI"), ("II", "2.IN-Software AI")):
        existing = (
            client.table("price_book_sections")
            .select("*")
            .eq("price_book_version_id", version_id)
            .eq("source_group_label", group_label)
            .limit(1)
            .execute()
        )
        if existing.data:
            section_ids[group_label] = existing.data[0]["id"]
        else:
            created = client.table("price_book_sections").insert(
                {
                    "price_book_version_id": version_id,
                    "source_sheet": source_sheet,
                    "source_group_label": group_label,
                    "display_label": group_label,
                    "name": f"Mục {group_label}",
                    "sort_order": 0 if group_label == "I" else 1,
                }
            ).execute()
            section_ids[group_label] = created.data[0]["id"]

    created_count = 0
    skipped_count = 0
    error_count = 0
    for item in all_items:
        group_label = "I" if item["source_sheet"] == "1.IN-Software AI" else "II"
        existing = (
            client.table("price_book_items")
            .select("id")
            .eq("price_book_version_id", version_id)
            .eq("source_sheet", item["source_sheet"])
            .eq("source_stt", item["source_stt"])
            .limit(1)
            .execute()
        )
        if existing.data:
            skipped_count += 1
            print(f"  BO QUA (da co): {item['sku']}")
            continue
        payload = {k: _decimal_to_jsonable(v) for k, v in item.items()}
        payload["price_book_version_id"] = version_id
        payload["section_id"] = section_ids[group_label]
        payload["status"] = "active"
        try:
            client.table("price_book_items").insert(payload).execute()
            created_count += 1
            print(f"  TAO MOI: {item['sku']}")
        except Exception as exc:  # noqa: BLE001
            error_count += 1
            print(f"  LOI: {item['sku']} -> {exc}")

    print(f"\nKET QUA: {created_count} tao moi / {skipped_count} bo qua (da co) / {error_count} loi")
    return 0 if error_count == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
