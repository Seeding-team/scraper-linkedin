"""Document engine hợp đồng: chỉnh trực tiếp DOCX giữ bố cục, điền bảng hạng mục, chốt chặn AI, PDF từ DOCX, phiên bản. Không cần mạng/DB."""
import asyncio
import base64
import hashlib
import io
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).parent))
from contract_docx_fixture import SAMPLE_DEAL, SAMPLE_QUOTE, build_sample_contract  # noqa: E402

from app.modules.all_platform.services import contract_ai_service as ai
from app.modules.all_platform.services import contract_docx_engine as eng
from app.modules.all_platform.services import contract_docx_pipeline as pipe


@pytest.fixture(scope="module")
def sample() -> bytes:
    return build_sample_contract()


def _doc(b):
    return eng.load_document(b)


def _find(doc, startswith):
    return next(p for p in eng.describe_paragraphs(doc) if p.text.startswith(startswith))


def run(coro):
    return asyncio.run(coro)


async def no_ai(paragraphs, deal, quote, prompt):
    return []


# ── giữ định dạng khi thay text ──
def test_replace_keeps_bold_label_and_run_formatting(sample):
    doc = _doc(sample)
    el = eng._all_paragraph_elements(doc)[int(_find(doc, "Mã số thuế").id[1:])]
    before_runs = [(r.find(qn("w:rPr")) is not None and r.find(qn("w:rPr")).find(qn("w:b")) is not None) for r in eng._text_runs(el)]
    assert eng.replace_paragraph_text(el, "Mã số thuế: 0312345678")
    runs = eng._text_runs(el)
    assert "".join(eng._run_text(r) for r in runs) == "Mã số thuế: 0312345678"
    assert eng._run_text(runs[0]) == "Mã số thuế: " and before_runs[0] is True      # nhãn vẫn nằm trong run in đậm gốc
    assert len(runs) == len(before_runs)                                             # không sinh/xoá run


def test_placeholders_filled_missing_kept_and_no_auto_date(sample):
    doc = _doc(sample)
    values = pipe.crm_values(SAMPLE_DEAL, SAMPLE_QUOTE, "HD-1", 205_700_000)      # không có sign_date
    assert "day" not in values and "year" not in values
    applied, missing = eng.fill_placeholders(doc, values)
    text = "\n".join(p.text for p in eng.describe_paragraphs(doc))
    assert "HD-1" in text and "CÔNG TY CỔ PHẦN GIẢI PHÁP ÁNH DƯƠNG" in text and "205.700.000" in text
    assert set(missing) == {"day", "month", "year"} and "{{day}}" in text            # ngày ký không tự lấy ngày hôm nay
    assert pipe.crm_values(None, None, sign_date="2026-10-09")["year"] == "2026"


# ── bảng hạng mục ──
def test_items_table_rows_cloned_with_prototype_formatting(sample):
    doc = _doc(sample)
    res = eng.fill_items_table(doc, SAMPLE_QUOTE)
    assert res["filled"] and res["rows"] == 3 and not res["warnings"]
    table = doc.tables[0]
    assert len(table.rows) == 1 + 4 + 1                                              # header + 3 hạng mục + 1 dòng nhóm (đúng thứ tự/số dòng báo giá) + tổng
    assert [r.cells[1].text for r in table.rows[1:5]] == ["Thiết kế bộ nhận diện thương hiệu", "Nhóm chiến dịch", "Chiến dịch quảng cáo đa kênh (3 tháng)", "Quản trị fanpage và sản xuất nội dung"]
    assert table.rows[3].cells[3].text == "38.000.000" and table.rows[3].cells[4].text == "125.400.000" and table.rows[3].cells[0].text == "2"   # thành tiền GỒM VAT
    assert table.rows[2].cells[0].text == "" and table.rows[2].cells[4].text == ""                                                              # dòng nhóm không có số
    assert table.rows[-1].cells[-1].text == "205.700.000"                            # tổng = totalAmount của báo giá
    for r in table.rows[1:5]:                                                        # shading/format của hàng mẫu được nhân bản
        assert r._tr.find(qn("w:tc")).find(qn("w:tcPr")).find(qn("w:shd")) is not None


def test_items_table_missing_or_no_quote_warns_instead_of_guessing(sample):
    res = eng.fill_items_table(_doc(sample), None)
    assert not res["filled"] and res["warnings"]
    from docx import Document
    plain = Document()
    plain.add_paragraph("Hợp đồng không có bảng")
    buf = io.BytesIO(); plain.save(buf)
    res2 = eng.fill_items_table(_doc(buf.getvalue()), SAMPLE_QUOTE)
    assert not res2["filled"] and "KHÔNG được chèn" in res2["warnings"][0]


# ── AI chỉ đề xuất, engine + chốt chặn áp ──
def test_validate_edits_blocks_invented_numbers_empty_and_unknown_ids():
    paras = {"p1": "3.1. Bên A thanh toán 100% ngay khi ký.", "p2": "Điều 4. Bảo hành."}
    ctx = "Thanh toán 50% khi ký, 40% khi bàn giao và 10% sau nghiệm thu. Thời gian triển khai 45 ngày."
    ok, bad = ai.validate_template_edits([
        {"id": "p1", "text": "3.1. Bên A thanh toán làm 03 đợt: đợt 1 là 50%, đợt 2 là 40%, đợt 3 là 10%."},
        {"id": "p2", "text": "Điều 4. Bảo hành 36 tháng và phạt 25% giá trị."},
        {"id": "p9", "text": "x"}, {"id": "p1", "text": "   "},
    ], paras, ctx)
    assert [e["id"] for e in ok] == ["p1"]
    reasons = " | ".join(b_["reason"] for b_ in bad)
    assert "25%" in reasons and "36 tháng" in reasons and "không tồn tại" in reasons and "xoá" in reasons


def test_pipeline_end_to_end_with_fake_ai_preserves_structure_and_original(sample):
    digest = hashlib.sha256(sample).hexdigest()

    async def ai_fake(paragraphs, deal, quote, prompt):
        pay = next(i for i, t in paragraphs if t.startswith("3.1."))
        return [{"id": pay, "text": "3.1. Bên A thanh toán làm 03 đợt: đợt 1 là 50%, đợt 2 là 40%, đợt 3 là 10%."}]

    res = run(pipe.render_from_template(sample, SAMPLE_DEAL, SAMPLE_QUOTE, "Thanh toán 50% khi ký, 40% khi bàn giao, 10% sau nghiệm thu.",
                                        "HD-2026-0042", 205_700_000, propose=ai_fake, sign_date="2026-10-09"))
    assert hashlib.sha256(sample).hexdigest() == digest                              # mẫu gốc không bị đổi
    assert res["structure"]["preserved"] and res["structure"]["rowCountChanged"] == [0]
    assert len(res["edits"]) == 1 and not res["rejectedEdits"] and not res["missingPlaceholders"]
    new = _doc(res["docx"])
    texts = [p.text for p in eng.describe_paragraphs(new)]
    assert any("đợt 1 là 50%" in t for t in texts) and any("ngày 09 tháng 10 năm 2026" in t for t in texts)
    # điều khoản không đụng tới giữ nguyên từng chữ
    old_texts = {p.text for p in eng.describe_paragraphs(_doc(sample)) if p.text.startswith(("4.", "5.", "6.", "7."))}
    assert old_texts <= set(texts)
    a, b = eng.fingerprint(_doc(sample)), eng.fingerprint(new)
    assert a["sections"] == b["sections"] and a["pageBreaks"] == b["pageBreaks"] == 1 and a["styles"] == b["styles"]


def test_pipeline_without_ai_or_context_does_not_call_ai(sample):
    called = []

    async def spy(*a):
        called.append(1)
        return []

    run(pipe.render_from_template(sample, None, None, None, propose=spy))
    assert not called


def test_ai_only_sees_non_table_paragraphs(sample):
    seen = {}

    async def spy(paragraphs, deal, quote, prompt):
        seen["ids"] = [i for i, _ in paragraphs]
        return []

    run(pipe.render_from_template(sample, SAMPLE_DEAL, SAMPLE_QUOTE, "x", propose=spy))
    infos = {p.id: p for p in eng.describe_paragraphs(_doc(sample))}
    assert seen["ids"] and not any(infos[i].in_table for i in seen["ids"])


# ── PDF ──
def test_pdf_requires_libreoffice_no_silent_fallback(monkeypatch, sample):
    monkeypatch.setattr(eng, "find_soffice", lambda: None)
    with pytest.raises(eng.PdfEngineUnavailable) as exc:
        eng.convert_docx_to_pdf(sample)
    assert "LibreOffice" in str(exc.value)


def _mini_pdf(content_stream: bytes, extra_catalog: bytes = b"", resources: bytes = b"<< /Font << /F1 4 0 R >> >>", extra_objs: bytes = b"") -> bytes:
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R " + extra_catalog + b">>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources " + resources + b" /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content_stream)).encode() + b" >>\nstream\n" + content_stream + b"\nendstream",
    ]
    if extra_objs:
        objs.append(extra_objs)
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode() + b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    return out + f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()


def test_pdf_template_classification_never_claims_layout_preserved():
    text_pdf = _mini_pdf(b"BT /F1 12 Tf 50 750 Td (" + b"Hop dong dich vu. " * 12 + b") Tj ET")
    form_pdf = _mini_pdf(b"BT /F1 12 Tf 50 750 Td (Hop dong) Tj ET", extra_catalog=b"/AcroForm << /Fields [] >> ")
    scan_pdf = _mini_pdf(b"q 500 0 0 700 50 50 cm /Im1 Do Q", resources=b"<< /XObject << /Im1 6 0 R >> >>",
                         extra_objs=b"<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceGray /BitsPerComponent 8 /Length 1 >>\nstream\n\x00\nendstream")
    t, f, s = (eng.classify_pdf_template(x) for x in (text_pdf, form_pdf, scan_pdf))
    assert (t["kind"], f["kind"], s["kind"]) == ("text", "form", "scan")
    assert not (t["layoutPreserved"] or f["layoutPreserved"] or s["layoutPreserved"])
    assert "DOCX" in t["warning"] and "scan" in s["warning"].lower() and "form" in f["warning"].lower()


@pytest.mark.skipif(eng.find_soffice() is None, reason="cần LibreOffice (chạy trong container contract-lo-test)")
def test_libreoffice_pdf_matches_docx_content(sample):
    res = run(pipe.render_from_template(sample, SAMPLE_DEAL, SAMPLE_QUOTE, None, "HD-1", 205_700_000, propose=no_ai, sign_date="2026-10-09"))
    pdf = eng.convert_docx_to_pdf(res["docx"])
    text = " ".join(eng.pdf_text(pdf).split())                  # dòng dài bị ngắt xuống dòng trong PDF -> chuẩn hoá khoảng trắng
    assert eng.pdf_page_count(pdf) >= 4
    for needle in ("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", "CÔNG TY CỔ PHẦN GIẢI PHÁP ÁNH DƯƠNG", "205.700.000", "ĐẠI DIỆN BÊN A", "HỢP ĐỒNG DỊCH VỤ"):
        assert needle in text
    assert "{{contract_number}}" not in text and "Trang" in text


# ── DOCX dựng từ điều khoản (không có mẫu) ──
def test_build_docx_from_clauses_has_standard_layout():
    docx = eng.build_docx_from_clauses("Hợp đồng dịch vụ", "HD-1", [{"title": "ĐIỀU 1. PHẠM VI", "body": "Dòng 1\nDòng 2"}], "Bên A", "Bên B")
    d = _doc(docx)
    texts = [p.text for p in eng.describe_paragraphs(d)]
    assert "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in texts and "ĐIỀU 1. PHẠM VI" in texts and "Số: HD-1" in texts
    assert d.sections[0].page_width.cm == pytest.approx(21, abs=0.01) and len(d.tables) == 1


# ── phiên bản / lưu trữ ──
class _Bucket:
    def __init__(self): self.files = {}
    def upload(self, key, content, file_options=None):
        if key in self.files:
            raise RuntimeError("exists")
        self.files[key] = content
    def download(self, key):
        if key not in self.files:
            raise KeyError(key)
        return self.files[key]
    def list(self, prefix):
        return [{"name": k.rsplit("/", 1)[1], "created_at": "t", "metadata": {"size": len(v)}} for k, v in self.files.items() if k.startswith(prefix + "/")]


def test_versions_never_overwrite_and_template_original_kept(monkeypatch, sample):
    from app.modules.all_platform.services import contract_document_store as store
    b = _Bucket()
    monkeypatch.setattr(store, "_bucket", lambda: b)
    assert store.save_version("C1", b"one") == 1 and store.save_version("C1", b"two") == 2
    assert [v["version"] for v in store.list_versions("C1")] == [1, 2]
    assert store.load_version("C1", 1) == b"one" and store.load_version("C1", 2) == b"two" and store.load_version("C1", 3) is None
    assert store.save_template_original("T1", "Mau.docx", sample)
    assert store.load_template_original("T1") == ("docx", sample)
    assert store.save_template_original("T1", "Mau.docx", b"ghi de") is None and store.load_template_original("T1")[1] == sample   # không ghi đè mẫu gốc


# ── router ──
def _client(monkeypatch, soffice=None):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.modules.all_platform.auth_deps import get_current_user
    from app.modules.all_platform.routers.contract_docx import contract_docx_router

    app = FastAPI()
    app.include_router(contract_docx_router, prefix="/contract-docs")
    app.dependency_overrides[get_current_user] = lambda: {"id": "u1"}
    monkeypatch.setattr(eng, "find_soffice", lambda: soffice)

    async def fake_chat(system_prompt, user_content):
        return {"edits": []}

    monkeypatch.setattr(ai, "_call_chat_json", fake_chat)
    return TestClient(app)


def test_render_endpoint_returns_docx_and_reports_pdf_unavailable_without_fallback(monkeypatch, sample):
    c = _client(monkeypatch)
    r = c.post("/contract-docs/render", files={"file": ("mau.docx", sample, "application/octet-stream")}, data={"extra_prompt": "điền thông tin", "contract_number": "HD-9"}).json()
    assert r["success"] and r["data"]["mode"] == "template-docx"
    assert base64.b64decode(r["data"]["docxBase64"])[:2] == b"PK"
    assert r["data"]["pdfBase64"] is None and "LibreOffice" in r["data"]["pdfError"]          # không có PDF 'thay thế'


def test_render_endpoint_pdf_template_is_flagged_not_pretended(monkeypatch):
    c = _client(monkeypatch)
    pdf = _mini_pdf(b"BT /F1 12 Tf 50 750 Td (" + b"Hop dong dich vu. " * 12 + b") Tj ET")
    r = c.post("/contract-docs/render", files={"file": ("mau.pdf", pdf, "application/pdf")}).json()
    assert r["success"] and r["data"]["mode"] == "pdf-reference" and r["data"]["layoutPreserved"] is False


def test_render_endpoint_rejects_non_docx_and_missing_template(monkeypatch):
    c = _client(monkeypatch)
    assert not c.post("/contract-docs/render", files={"file": ("mau.txt", b"abc", "text/plain")}).json()["success"]
    assert "Chưa chọn mẫu" in c.post("/contract-docs/render", data={}).json()["message"]


def test_from_clauses_endpoint(monkeypatch):
    c = _client(monkeypatch)
    r = c.post("/contract-docs/from-clauses", json={"title": "HĐ", "contract_number": "1", "clauses": [{"title": "ĐIỀU 1", "body": "x"}]}).json()
    assert r["success"] and r["data"]["mode"] == "generated-docx" and r["data"]["pdfBase64"] is None


def test_placeholder_split_across_runs_like_word_does():
    from docx import Document
    d = Document()
    p = d.add_paragraph()
    p.add_run("Số hợp đồng: ")
    p.add_run("{{con").bold = True
    p.add_run("tract_")
    p.add_run("number}}").italic = True
    p.add_run(" (bản chính)")
    buf = io.BytesIO(); d.save(buf)
    doc = _doc(buf.getvalue())
    applied, missing = eng.fill_placeholders(doc, {"contract_number": "HD-7"})
    assert applied and not missing
    para = doc.paragraphs[0]
    assert para.text == "Số hợp đồng: HD-7 (bản chính)"
    assert para.runs[0].text == "Số hợp đồng: " and para.runs[-1].text == " (bản chính)"


# ───────── bug thật phát hiện trên file mẫu thật: hàng TỔNG CỘNG của bảng hạng mục GỘP Ô nhãn (vd "TỔNG CỘNG" chiếm
# 2 cột đầu qua gridSpan) nên <w:tc> CUỐI CÙNG trong XML không phải cột "Phí dịch vụ" (amount) mà lại là cột "Đơn vị"
# phía sau nó - code cũ giả định "ô số tiền luôn là ô cuối" nên ghi NHẦM tổng tiền vào cột Đơn vị, cột tiền giữ số cũ
# của mẫu (vd vẫn "5.000.000" dù báo giá thật là 111.375.177) ─────────
def test_items_table_total_row_with_merged_label_cell_updates_correct_column():
    from docx import Document

    doc = Document()
    t = doc.add_table(rows=3, cols=4)
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "STT", "Hạng mục", "Phí dịch vụ", "Đơn vị"
    row = t.rows[1].cells
    row[0].text, row[1].text, row[2].text, row[3].text = "1", "Phí sử dụng Nền tảng", "5.000.000", "VNĐ"
    total_row = t.rows[2].cells
    total_row[0].merge(total_row[1])                 # giống mẫu thật: nhãn "TỔNG CỘNG" gộp 2 cột đầu
    total_row[0].text = "TỔNG CỘNG (trọn gói 01 tháng)"
    total_row[2].text = "5.000.000"
    total_row[3].text = "VNĐ"
    buf = io.BytesIO(); doc.save(buf)
    doc2 = _doc(buf.getvalue())

    quote = {"currency": "VND", "subtotalAmount": 103125177, "vatAmount": 8250000, "totalAmount": 111375177,
             "items": [{"description": "test", "unit": "test", "quantity": 1, "unitPrice": 159, "vatRate": 0, "totalAmount": 159}]}
    res = eng.fill_items_table(doc2, quote)
    assert res["filled"] and not res["warnings"]
    t2 = doc2.tables[0]
    last_row_cells = [c.text for c in t2.rows[-1].cells]
    assert "111.375.177" in last_row_cells           # tổng tiền đúng, không còn "5.000.000" cũ
    assert "5.000.000" not in last_row_cells
    assert last_row_cells[-1] == "VNĐ"               # cột Đơn vị KHÔNG bị ghi đè nhầm số tiền vào


# ───────── "sửa tự do" 1 hợp đồng ĐÃ TẠO: thêm/bớt đoạn tùy ý (khác apply_edits() chỉ thay text đoạn ĐÃ CÓ) ─────────
def _body_texts(doc) -> list[str]:
    return [p.text for p in doc.paragraphs if p.text.strip()]


def test_replace_body_paragraphs_update_only_keeps_same_count():
    from docx import Document

    doc = Document()
    doc.add_paragraph("ĐIỀU 1. Mục đích")
    doc.add_paragraph("Nội dung điều 1.")
    doc.add_paragraph("ĐIỀU 2. Thanh toán")
    buf = io.BytesIO(); doc.save(buf)
    doc2 = _doc(buf.getvalue())
    changes = eng.replace_body_paragraphs(doc2, ["ĐIỀU 1. Mục đích", "Nội dung điều 1 đã sửa.", "ĐIỀU 2. Thanh toán"])
    assert len(changes) == 1 and changes[0]["op"] == "update"
    assert _body_texts(doc2) == ["ĐIỀU 1. Mục đích", "Nội dung điều 1 đã sửa.", "ĐIỀU 2. Thanh toán"]


def test_replace_body_paragraphs_can_add_new_lines_no_limit():
    from docx import Document

    doc = Document()
    doc.add_paragraph("ĐIỀU 1. Mục đích")
    doc.add_paragraph("Nội dung điều 1.")
    buf = io.BytesIO(); doc.save(buf)
    doc2 = _doc(buf.getvalue())
    new_lines = ["ĐIỀU 1. Mục đích", "Nội dung điều 1.", "ĐIỀU 2. Thêm mới", "2.1. Ý đầu.", "2.2. Ý sau."]
    changes = eng.replace_body_paragraphs(doc2, new_lines)
    assert sum(1 for c in changes if c["op"] == "insert") == 3
    assert _body_texts(doc2) == new_lines


def test_replace_body_paragraphs_can_delete_lines():
    from docx import Document

    doc = Document()
    doc.add_paragraph("ĐIỀU 1.")
    doc.add_paragraph("Dòng sẽ bị xoá.")
    doc.add_paragraph("ĐIỀU 2.")
    buf = io.BytesIO(); doc.save(buf)
    doc2 = _doc(buf.getvalue())
    # Diff theo VỊ TRÍ (không phải LCS thông minh) - xoá 1 dòng giữa có thể ghi nhận thành "sửa dòng kế + xoá dòng
    # cuối" thay vì "xoá đúng dòng giữa", nhưng KẾT QUẢ CUỐI vẫn đúng - đó là điều thực sự quan trọng.
    changes = eng.replace_body_paragraphs(doc2, ["ĐIỀU 1.", "ĐIỀU 2."])
    assert len(changes) >= 1
    assert "Dòng sẽ bị xoá." not in _body_texts(doc2)
    assert _body_texts(doc2) == ["ĐIỀU 1.", "ĐIỀU 2."]


def test_replace_body_paragraphs_does_not_touch_tables():
    from docx import Document

    doc = Document()
    doc.add_paragraph("ĐIỀU 1. Mục đích")
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "STT", "Hạng mục"
    buf = io.BytesIO(); doc.save(buf)
    doc2 = _doc(buf.getvalue())
    eng.replace_body_paragraphs(doc2, ["ĐIỀU 1. Mục đích mới", "ĐIỀU 2. Thêm"])
    assert doc2.tables[0].rows[0].cells[0].text == "STT"          # bảng giữ nguyên hoàn toàn
    assert doc2.tables[0].rows[0].cells[1].text == "Hạng mục"


def test_replace_body_endpoint(monkeypatch):
    c = _client(monkeypatch)
    from docx import Document

    d = Document()
    d.add_paragraph("ĐIỀU 1. Mục đích")
    d.add_paragraph("Nội dung cũ.")
    buf = io.BytesIO(); d.save(buf)
    docx_b64 = base64.b64encode(buf.getvalue()).decode()
    r = c.post("/contract-docs/replace-body", json={"docx_base64": docx_b64, "lines": ["ĐIỀU 1. Mục đích", "Nội dung mới.", "ĐIỀU 2. Thêm mới"]}).json()
    assert r["success"], r
    assert r["data"]["changes"] and len(r["data"]["changes"]) == 2
    out_doc = _doc(base64.b64decode(r["data"]["docxBase64"]))
    assert _body_texts(out_doc) == ["ĐIỀU 1. Mục đích", "Nội dung mới.", "ĐIỀU 2. Thêm mới"]
