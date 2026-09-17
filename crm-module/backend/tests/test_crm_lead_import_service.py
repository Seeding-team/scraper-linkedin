from __future__ import annotations

import io
import importlib.util
import sys
import types
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook


def _module(name: str, **attributes):
    module = types.ModuleType(name)
    for key, value in attributes.items():
        setattr(module, key, value)
    sys.modules[name] = module
    return module


for package_name in ("app", "app.core", "app.modules", "app.modules.all_platform", "app.modules.all_platform.services"):
    package = _module(package_name)
    package.__path__ = []

_module("app.core.config", settings=SimpleNamespace(crm_instance="markee"))
_module("app.core.supabase_client", execute_supabase_query=lambda callback: callback(), get_supabase_client=lambda: None)


class _DuplicateLeadError(ValueError):
    pass


_module(
    "app.modules.all_platform.services.crm_lead_service",
    DuplicateLeadError=_DuplicateLeadError,
    create_lead=lambda _payload, _user: {},
    is_valid_email=lambda value: bool(__import__("re").fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", str(value or "").strip())),
)


def _normalize_email(value):
    return str(value).strip().lower() if value else None


def _normalize_phone(value):
    text = "".join(ch for ch in str(value or "") if ch.isdigit())
    if len(text) == 10 and text.startswith("0"):
        return "+84" + text[1:]
    if len(text) == 11 and text.startswith("84"):
        return "+" + text
    return None


_module(
    "app.modules.all_platform.services.crm_customer_service",
    normalize_email=_normalize_email,
    normalize_phone=_normalize_phone,
)
_module(
    "app.modules.all_platform.services.crm_permission_service",
    has_full_crm_access=lambda _user: True,
    can_manage_shared_master_data=lambda _user: False,
)
_module(
    "app.modules.all_platform.services.supabase_categories_service",
    add_category=lambda payload: {**payload, "id": f"new-{payload.get('code')}"},
)

_SERVICE_PATH = Path(__file__).parents[1] / "app/modules/all_platform/services/crm_lead_import_service.py"
_SPEC = importlib.util.spec_from_file_location("crm_lead_import_service_under_test", _SERVICE_PATH)
assert _SPEC and _SPEC.loader
service = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(service)


class _Result:
    def __init__(self, data):
        self.data = data


class _EmptyQuery:
    @property
    def not_(self): return self
    def select(self, *_args): return self
    def eq(self, *_args): return self
    def in_(self, *_args): return self
    def is_(self, *_args): return self
    def execute(self): return _Result([])


class _EmptySupabase:
    def table(self, _name): return _EmptyQuery()


def _xlsx(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(service.FIELD_HEADERS.values()))
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_preview_marks_valid_database_duplicate_file_duplicate_and_bad_owner(monkeypatch):
    categories = {
        "crm_source": [{"id": "source-1", "code": "website", "name": "Website"}],
        "crm_position": [{"id": "position-1", "code": "director", "name": "Giám đốc"}],
    }
    monkeypatch.setattr(service, "_tenant_categories", lambda kind: categories[kind])
    monkeypatch.setattr(service, "_active_users", lambda: [
        {"id": "owner-1", "name": "Sale A", "email": "sale@example.com", "is_active": True}
    ])
    monkeypatch.setattr(service, "has_full_crm_access", lambda _user: True)

    filters = []

    class _Query:
        @property
        def not_(self): return self
        def select(self, *_args): return self
        def eq(self, *args): filters.append(args); return self
        def in_(self, *_args): return self
        def is_(self, *_args): return self
        def execute(self):
            return _Result([{
                "id": "lead-old", "lead_name": "Nguyễn Văn Cũ",
                "phone_normalized": "+84901111111", "email_normalized": None,
            }])

    class _Supabase:
        def table(self, _name): return _Query()

    monkeypatch.setattr(service, "get_supabase_client", lambda: _Supabase())
    monkeypatch.setattr(service, "execute_supabase_query", lambda callback: callback())
    raw = _xlsx([
        ["Lead A", "Công ty A", "0902222222", "a@example.com", "Giám đốc", "Website", "Sale A", None, None, None, None, None],
        ["Lead cũ", "Công ty B", "0901111111", None, None, "Website", "Sale A", None, None, None, None, None],
        ["Lead trùng file", "Công ty C", "0902222222", None, None, "Website", "Sale A", None, None, None, None, None],
        ["Lead owner lỗi", "Công ty D", "0903333333", None, None, "Website", "Không tồn tại", None, None, None, None, None],
    ])

    preview = service.preview_import(raw, {"id": "owner-1"})

    assert preview["summary"] == {"total": 4, "valid": 1, "duplicate": 2, "error": 1}
    assert preview["rows"][0]["data"]["position_category_id"] == "position-1"
    assert preview["rows"][0]["data"]["source"] == "website"
    assert preview["rows"][3]["issues"][0]["column"] == "Người phụ trách Lead"
    assert ("instance", "markee") in filters

    dup_issue = next(issue for issue in preview["rows"][1]["issues"] if issue["code"] == "duplicate_database")
    assert dup_issue["lead_id"] == "lead-old"
    assert dup_issue["lead_name"] == "Nguyễn Văn Cũ"
    assert "Nguyễn Văn Cũ" in dup_issue["message"]


def test_categories_are_shared_active_master_data_without_instance_column(monkeypatch):
    filters = []

    class Query:
        def select(self, columns):
            assert 'instance' not in columns
            return self
        def eq(self, *args):
            filters.append(args)
            return self
        def execute(self):
            assert filters == [('category_type', 'crm_source'), ('is_active', True)]
            return _Result([{'id': 'shared-source', 'code': 'Existing_Customer', 'name': 'Shared source'}])

    class Supabase:
        def table(self, name):
            assert name == 'categories'
            return Query()

    monkeypatch.setattr(service, 'get_supabase_client', lambda: Supabase())
    monkeypatch.setattr(service, 'execute_supabase_query', lambda callback: callback())
    monkeypatch.setattr(service, 'settings', SimpleNamespace(crm_instance='SECURITYZONE'))
    assert service._tenant_categories('crm_source')[0]['code'] == 'Existing_Customer'


def test_pending_create_when_source_unmatched_and_user_can_manage_master_data(monkeypatch):
    monkeypatch.setattr(service, "_tenant_categories", lambda _kind: [])
    monkeypatch.setattr(service, "_active_users", lambda: [])
    monkeypatch.setattr(service, "has_full_crm_access", lambda _user: True)
    monkeypatch.setattr(service, "get_supabase_client", lambda: _EmptySupabase())
    monkeypatch.setattr(service, "execute_supabase_query", lambda callback: callback())

    parsed = [(2, {
        "lead_name": "A", "phone": "0902222222", "email": None,
        "source": "TikTok Ads Q4", "position": None, "owner": None,
    })]
    result = service._validate_and_resolve_rows(parsed, {"id": "owner-1"}, allow_create_master=True)
    row = result["rows"][0]

    assert row["status"] == "valid"
    assert row["pending_creates"] == [{"field": "source", "value": "TikTok Ads Q4"}]
    # Not resolved to a real category id/code yet — that only happens at
    # confirm time via _resolve_or_create_pending_masters.
    assert row["data"]["source"] == "TikTok Ads Q4"


def test_error_when_source_unmatched_and_user_cannot_manage_master_data(monkeypatch):
    monkeypatch.setattr(service, "_tenant_categories", lambda _kind: [])
    monkeypatch.setattr(service, "_active_users", lambda: [])
    monkeypatch.setattr(service, "has_full_crm_access", lambda _user: True)
    monkeypatch.setattr(service, "get_supabase_client", lambda: _EmptySupabase())
    monkeypatch.setattr(service, "execute_supabase_query", lambda callback: callback())

    parsed = [(2, {
        "lead_name": "A", "phone": "0902222222", "email": None,
        "source": "TikTok Ads Q4", "position": None, "owner": None,
    })]
    result = service._validate_and_resolve_rows(parsed, {"id": "owner-1"}, allow_create_master=False)
    row = result["rows"][0]

    assert row["status"] == "error"
    assert row["pending_creates"] == []
    issue = next(issue for issue in row["issues"] if issue["code"] == "unknown_source")
    assert "TikTok Ads Q4" in issue["message"]
    assert "không có quyền" in issue["message"]


def test_owner_forbidden_when_edited_row_assigns_to_other_without_permission(monkeypatch):
    monkeypatch.setattr(service, "_tenant_categories", lambda _kind: [])
    monkeypatch.setattr(service, "_active_users", lambda: [
        {"id": "owner-2", "name": "Sale B", "email": "b@example.com", "is_active": True},
    ])
    monkeypatch.setattr(service, "has_full_crm_access", lambda _user: False)
    monkeypatch.setattr(service, "get_supabase_client", lambda: _EmptySupabase())
    monkeypatch.setattr(service, "execute_supabase_query", lambda callback: callback())

    # FE member combobox sends the picked member's id as "owner" text (same
    # field Excel-typed names/emails use) — the "id" lookup key resolves it.
    parsed = [(2, {
        "lead_name": "A", "phone": "0902222222", "email": None,
        "source": None, "position": None, "owner": "owner-2",
    })]
    result = service._validate_and_resolve_rows(parsed, {"id": "owner-1"}, allow_create_master=False)
    row = result["rows"][0]

    assert row["status"] == "error"
    issue = next(issue for issue in row["issues"] if issue["code"] == "owner_forbidden")
    assert "Sale B" in issue["message"]


def test_resolve_or_create_pending_masters_dedupes_new_source_across_batch(monkeypatch):
    monkeypatch.setattr(service, "can_manage_shared_master_data", lambda _user: True)
    monkeypatch.setattr(service, "_tenant_categories", lambda _kind: [])
    calls = []

    def _add_category(payload):
        calls.append(payload)
        return {"id": "new-source-id", "code": payload["code"], "name": payload["name"]}

    monkeypatch.setattr(service, "add_category", _add_category)
    rows = [
        {"data": {"source": None}, "pending_creates": [{"field": "source", "value": "TikTok Ads Q4"}]},
        {"data": {"source": None}, "pending_creates": [{"field": "source", "value": "tiktok ads q4"}]},
        {"data": {"source": None}, "pending_creates": [{"field": "source", "value": "TikTok Ads Q4"}]},
    ]

    service._resolve_or_create_pending_masters(rows, {"id": "owner-1"})

    assert len(calls) == 1
    assert all(row["data"]["source"] == "TikTok Ads Q4" for row in rows)
    assert all(row["pending_creates"] == [] for row in rows)


def test_confirm_revalidates_and_only_creates_selected_valid_rows(monkeypatch):
    revalidated = {
        "rows": [
            {
                "row_number": 2, "status": "valid",
                "data": {"lead_name": "A", "phone": "0902222222", "owner_name": "Sale A", "sdr_id": "owner-1"},
                "issues": [], "pending_creates": [],
            },
            {
                "row_number": 3, "status": "duplicate",
                "data": {"lead_name": "B"}, "issues": [], "pending_creates": [],
            },
        ],
    }
    monkeypatch.setattr(service, "revalidate_rows", lambda _rows, _user: revalidated)
    created_payloads = []

    def _create(payload, _user):
        created_payloads.append(payload)
        return {"id": "new-lead"}

    monkeypatch.setattr(service, "create_lead", _create)
    result = service.confirm_import([{"row_number": 2}, {"row_number": 3}], [2, 3], {"id": "owner-1"})

    assert result["created"] == 1
    assert result["skipped"] == 1
    assert created_payloads == [{"lead_name": "A", "phone": "0902222222", "sdr_id": "owner-1", "status": "mql"}]


def test_template_contains_current_master_data_and_safe_excel_values(monkeypatch):
    categories = {
        "crm_source": [{"id": "source-1", "code": "website", "name": "Website"}],
        "crm_position": [{"id": "position-1", "code": "director", "name": "Giám đốc"}],
    }
    monkeypatch.setattr(service, "_tenant_categories", lambda kind: categories[kind])
    monkeypatch.setattr(service, "_active_users", lambda: [
        {"id": "owner-1", "name": "=unsafe", "email": "sale@example.com", "is_active": True}
    ])

    workbook = load_workbook(io.BytesIO(service.build_template()))

    assert workbook["Leads"]["A1"].value == "Họ tên người liên hệ"
    assert workbook["Leads"]["C2"].number_format == "@"
    assert workbook["Danh mục tham chiếu"].sheet_state == "hidden"
    assert workbook["Danh mục tham chiếu"]["D2"].value == "'=unsafe"
