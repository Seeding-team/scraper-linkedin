"""Dựng HỢP ĐỒNG MẪU DOCX thực tế (quốc hiệu, placeholder, bảng hạng mục, nhiều điều khoản nhiều trang, header/footer có số trang,
khu vực chữ ký, page break) để test document engine. Không phụ thuộc app/* nên chạy được cả trong container LibreOffice."""
from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"


def _font(run, size=13, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold, run.italic = bold, italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)


def _para(doc, label="", value="", align=None, bold_value=False, size=13, space_after=6, keep_bold_label=True):
    p = doc.add_paragraph()
    if label:
        _font(p.add_run(label), size, bold=keep_bold_label)
    if value:
        _font(p.add_run(value), size, bold=bold_value)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    return p


def _page_field(paragraph):
    run = paragraph.add_run()
    _font(run, 10)
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            el.text = text
        run._element.append(el)


def _shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


CLAUSE_FILLER = (
    "Hai bên cam kết thực hiện đầy đủ các nghĩa vụ đã thỏa thuận trong hợp đồng này theo nguyên tắc thiện chí, hợp tác và trung thực. "
    "Mọi thay đổi, bổ sung phải được lập thành văn bản và có chữ ký xác nhận của đại diện hợp pháp của cả hai bên trước khi có hiệu lực. "
    "Trong trường hợp phát sinh tranh chấp, hai bên ưu tiên giải quyết thông qua thương lượng; nếu không đạt được thỏa thuận thì đưa ra Tòa án có thẩm quyền."
)


def build_sample_contract() -> bytes:
    import io

    doc = Document()
    normal = doc.styles["Normal"]                      # font mặc định tường minh (không để LibreOffice rơi về font theme khác)
    normal.font.name = FONT
    nrpr = normal.element.get_or_add_rPr()
    nrf = nrpr.find(qn("w:rFonts"))
    if nrf is None:
        nrf = OxmlElement("w:rFonts")
        nrpr.insert(0, nrf)
    for a_ in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        nrf.set(qn(a_), FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(2), Cm(2.5), Cm(2.5)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(hp.add_run("CÔNG TY TNHH MARKEE — HỢP ĐỒNG DỊCH VỤ"), 9, italic=True)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(fp.add_run("Trang "), 10)
    _page_field(fp)

    C = WD_ALIGN_PARAGRAPH.CENTER
    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    _para(doc, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", align=C, space_after=0)
    doc.paragraphs[-1].runs[0].font.size = Pt(13)
    _para(doc, "Độc lập - Tự do - Hạnh phúc", align=C, space_after=2)
    _para(doc, "—————", align=C, space_after=10, keep_bold_label=False)
    _para(doc, "HỢP ĐỒNG CUNG CẤP DỊCH VỤ", align=C, space_after=2).runs[0].font.size = Pt(15)
    _para(doc, "Số: ", "{{contract_number}}", align=C, space_after=10)
    _para(doc, "", "Hôm nay, ngày {{day}} tháng {{month}} năm {{year}}, tại Thành phố Hồ Chí Minh, chúng tôi gồm:", align=J)
    _para(doc, "BÊN A (Bên sử dụng dịch vụ): ", "{{company_name}}", bold_value=True)
    _para(doc, "Mã số thuế: ", "{{tax_code}}")
    _para(doc, "Địa chỉ: ", "{{address}}")
    _para(doc, "Người đại diện: ", "{{customer_name}} — Chức vụ: {{representative_position}}")
    _para(doc, "BÊN B (Bên cung cấp dịch vụ): ", "CÔNG TY TNHH MARKEE", bold_value=True)
    _para(doc, "Địa chỉ: ", "Số 1 Đường Nguyễn Huệ, Quận 1, TP. Hồ Chí Minh")
    _para(doc, "", "Hai bên cùng thống nhất ký kết hợp đồng với các điều khoản sau:", align=J)

    _para(doc, "ĐIỀU 1. PHẠM VI CÔNG VIỆC", space_after=3)
    p = _para(doc, "", "Bên B cung cấp cho Bên A các dịch vụ được liệt kê trong Bảng hạng mục dưới đây, theo đúng yêu cầu kỹ thuật đã thống nhất.", align=J)
    p.paragraph_format.left_indent = Cm(0.5)

    _para(doc, "ĐIỀU 2. BẢNG HẠNG MỤC VÀ GIÁ TRỊ HỢP ĐỒNG", space_after=3)
    table = doc.add_table(rows=3, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(1.2), Cm(6.3), Cm(1.8), Cm(3.0), Cm(3.0)]
    heads = ["STT", "Hạng mục", "Số lượng", "Đơn giá (VNĐ)", "Thành tiền (VNĐ)"]
    for ci, h in enumerate(heads):
        c = table.cell(0, ci)
        c.width = widths[ci]
        c.text = ""
        _font(c.paragraphs[0].add_run(h), 11, bold=True)
        c.paragraphs[0].alignment = C
        _shade(c, "D9E2F3")
    sample = ["1", "Hạng mục mẫu", "1", "0", "0"]
    for ci, v in enumerate(sample):
        c = table.cell(1, ci)
        c.width = widths[ci]
        c.text = ""
        _font(c.paragraphs[0].add_run(v), 11)
        c.paragraphs[0].alignment = C if ci in (0, 2) else (WD_ALIGN_PARAGRAPH.RIGHT if ci > 2 else WD_ALIGN_PARAGRAPH.LEFT)
        _shade(c, "F2F2F2")
    tot = table.rows[2]
    tot.cells[0].merge(tot.cells[3])
    tot.cells[0].text = ""
    _font(tot.cells[0].paragraphs[0].add_run("Tổng cộng"), 11, bold=True)
    tot.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tot.cells[4].text = ""
    _font(tot.cells[4].paragraphs[0].add_run("0"), 11, bold=True)
    tot.cells[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _para(doc, "", "Giá trị hợp đồng: {{contract_value}} VNĐ (đã bao gồm thuế GTGT theo báo giá {{quote_number}}).", align=J, space_after=8)

    _para(doc, "ĐIỀU 3. THANH TOÁN", space_after=3)
    _para(doc, "", "3.1. Bên A thanh toán 100% giá trị hợp đồng ngay sau khi ký kết hợp đồng này.", align=J)
    _para(doc, "", "3.2. Hình thức thanh toán: chuyển khoản vào tài khoản của Bên B.", align=J)
    _para(doc, "", "3.3. Thời gian triển khai dự kiến 30 ngày kể từ ngày hợp đồng có hiệu lực.", align=J)

    for n, name in ((4, "TRIỂN KHAI VÀ NGHIỆM THU"), (5, "BẢO MẬT THÔNG TIN"), (6, "TRÁCH NHIỆM VÀ PHẠT VI PHẠM")):
        _para(doc, f"ĐIỀU {n}. {name}", space_after=3)
        for k in range(1, 6):
            _para(doc, "", f"{n}.{k}. {CLAUSE_FILLER}", align=J)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _para(doc, "ĐIỀU 7. ĐIỀU KHOẢN CHUNG", space_after=3)
    for k in range(1, 9):
        _para(doc, "", f"7.{k}. {CLAUSE_FILLER}", align=J)
    _para(doc, "", "Hợp đồng được lập thành 02 bản có giá trị pháp lý như nhau, mỗi bên giữ 01 bản.", align=J, space_after=14)

    sig = doc.add_table(rows=2, cols=2)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ci, who in enumerate(("ĐẠI DIỆN BÊN A", "ĐẠI DIỆN BÊN B")):
        c = sig.cell(0, ci)
        c.text = ""
        _font(c.paragraphs[0].add_run(who), 13, bold=True)
        c.paragraphs[0].alignment = C
        d = sig.cell(1, ci)
        d.text = ""
        _font(d.paragraphs[0].add_run("(Ký, ghi rõ họ tên, đóng dấu)"), 11, italic=True)
        d.paragraphs[0].alignment = C
        for _ in range(4):
            d.add_paragraph()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


SAMPLE_DEAL = {
    "customer_name": "Nguyễn Văn An", "company_name": "CÔNG TY CỔ PHẦN GIẢI PHÁP ÁNH DƯƠNG", "tax_code": "0312345678",
    "address": "Số 25 Lê Lợi, Phường Bến Nghé, Quận 1, TP. Hồ Chí Minh", "position": "Giám đốc", "email": "an@anhduong.vn", "phone": "0901234567",
}
SAMPLE_QUOTE = {
    "quoteNumber": "BG-2026-0042", "currency": "VND", "subtotalAmount": 187_000_000, "vatAmount": 18_700_000, "totalAmount": 205_700_000,
    "items": [
        {"rowType": "item", "description": "Thiết kế bộ nhận diện thương hiệu", "quantity": 1, "unitPrice": 45_000_000, "amountAfterDiscount": 45_000_000, "vatRate": 10, "totalAmount": 49_500_000, "serviceDescription": "Logo, bảng màu, sổ tay thương hiệu", "unit": "Gói"},
        {"rowType": "section", "description": "Nhóm chiến dịch"},
        {"rowType": "item", "description": "Chiến dịch quảng cáo đa kênh (3 tháng)", "quantity": 3, "unitPrice": 38_000_000, "amountAfterDiscount": 114_000_000, "vatRate": 10, "totalAmount": 125_400_000, "serviceDescription": "Facebook, Google, TikTok", "unit": "Tháng"},
        {"rowType": "item", "description": "Quản trị fanpage và sản xuất nội dung", "quantity": 4, "unitPrice": 7_000_000, "amountAfterDiscount": 28_000_000, "vatRate": 10, "totalAmount": 30_800_000, "serviceDescription": "", "unit": "Tháng"},
    ],
}
