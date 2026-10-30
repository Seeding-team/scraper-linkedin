"""Unit tests for Task 3: DOCX -> PDF formatting fixes (3-col party table, cantSplit, keepNext, totals_kind, find_items_table)."""
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.modules.all_platform.services import contract_docx_engine as eng


def _make_party_doc():
    doc = Document()
    # Paragraph marking BÊN A
    doc.add_paragraph("BÊN A (BÊN THUÊ DỊCH VỤ): CÔNG TY CỔ PHẦN NÔNG NGHIỆP U&I")
    tbl_a = doc.add_table(rows=4, cols=3)
    # Row 0: 3 columns [Label | : | Value]
    r0 = tbl_a.rows[0]
    r0.cells[0].text = "Tên công ty"
    r0.cells[1].text = ":"
    r0.cells[2].text = "CÔNG TY CỔ PHẦN CŨ"
    # Row 1: 3 columns with empty middle or colon
    r1 = tbl_a.rows[1]
    r1.cells[0].text = "Mã số thuế"
    r1.cells[1].text = ":"
    r1.cells[2].text = "0100000000"
    # Row 2: 3 columns with empty value cell (no runs)
    r2 = tbl_a.rows[2]
    r2.cells[0].text = "Điện thoại"
    r2.cells[1].text = ":"
    r2.cells[2].text = ""  # empty
    # Row 3: 2-column inline merged row where cell 0 has "Địa chỉ : ..." and cell 1 is empty narrow cell
    r3 = tbl_a.rows[3]
    r3.cells[0].text = "Địa chỉ                    : Số 123 Đường Cũ"
    r3.cells[1].text = ""
    # Remove cell 2 from row 3 so len(tcs) == 2
    r3._tr.remove(r3.cells[2]._tc)

    return doc, tbl_a


def test_three_column_party_table_fills_col2_not_colon_col1():
    doc, tbl = _make_party_doc()
    party_a = {
        "tax_code": "3701528386",
        "phone": "02743831111",
        "address": "Số 9 Ngô Gia Tự, Bình Dương",
    }
    parties = {"a": party_a}
    res = eng.fill_party_info_tables(doc, parties)

    assert len(res["filled"]) >= 3
    # Verify Row 1 (MST): cell 1 must remain ":" while cell 2 gets "3701528386"
    assert tbl.rows[1].cells[1].text.strip() == ":"
    assert tbl.rows[1].cells[2].text.strip() == "3701528386"

    # Verify Row 2 (Phone, originally empty): cell 1 remains ":" while cell 2 gets "02743831111"
    assert tbl.rows[2].cells[1].text.strip() == ":"
    assert tbl.rows[2].cells[2].text.strip() == "02743831111"

    # Verify Row 3 (Inline Địa chỉ in cell 0): cell 0 gets updated after ":" and cell 1 remains empty
    assert tbl.rows[3].cells[1].text.strip() == ""
    assert "Số 9 Ngô Gia Tự, Bình Dương" in tbl.rows[3].cells[0].text
    assert tbl.rows[3].cells[0].text.startswith("Địa chỉ")


def test_totals_kind_distinguishes_tax_exempt_and_part_subtotals():
    # Tax-exempt must NOT be matched as VAT
    assert eng.totals_kind("Thuế GTGT: Không chịu thuế") is None
    assert eng.totals_kind("Không chịu thuế GTGT") is None
    assert eng.totals_kind("Thuế GTGT (miễn thuế)") is None

    # Part subtotals must NOT be matched as grand total
    assert eng.totals_kind("Cộng Phần A (chưa giảm giá)") == "subtotal_before"
    assert eng.totals_kind("Cộng Phần A") == "part_subtotal"
    assert eng.totals_kind("Cộng Phần B") == "part_subtotal"
    assert eng.totals_kind("Cộng Gói 1") == "part_subtotal"

    # Real totals
    assert eng.totals_kind("Cộng tiền hàng (chưa VAT)") == "subtotal"
    assert eng.totals_kind("Thuế GTGT 10%") == "vat"
    assert eng.totals_kind("TỔNG GIÁ TRỊ HỢP ĐỒNG (A + B)") == "total"
    assert eng.totals_kind("Tổng cộng thanh toán") == "total"


def test_find_items_table_prefers_detail_table_over_summary_table():
    doc = Document()
    # Table 0: Summary table [STT | Nội dung | Thành tiền] (3 cols)
    t_sum = doc.add_table(rows=3, cols=3)
    t_sum.rows[0].cells[0].text = "STT"
    t_sum.rows[0].cells[1].text = "Nội dung"
    t_sum.rows[0].cells[2].text = "Thành tiền (VNĐ)"

    # Table 1: Detail table [STT | Hạng mục | ĐVT | Số lượng | Đơn giá | Thành tiền] (6 cols)
    t_det = doc.add_table(rows=3, cols=6)
    t_det.rows[0].cells[0].text = "STT"
    t_det.rows[0].cells[1].text = "Hạng mục"
    t_det.rows[0].cells[2].text = "ĐVT"
    t_det.rows[0].cells[3].text = "Số lượng"
    t_det.rows[0].cells[4].text = "Đơn giá"
    t_det.rows[0].cells[5].text = "Thành tiền"

    found = eng.find_items_table(doc)
    assert found is not None
    table, h_idx, cols = found
    # Must select Table 1 (the detail table), not Table 0!
    assert table._tbl is t_det._tbl
    assert "unit_price" in cols
    assert "qty" in cols


def test_table_row_splits_and_keep_next_prevention():
    doc = Document()
    tbl = doc.add_table(rows=5, cols=3)
    tbl.rows[0].cells[0].text = "STT"
    tbl.rows[0].cells[1].text = "Nội dung"
    tbl.rows[0].cells[2].text = "Thành tiền"

    tbl.rows[1].cells[0].text = "1"
    tbl.rows[1].cells[1].text = "Phần mềm CloudGate"
    tbl.rows[1].cells[2].text = "27.272.724"

    tbl.rows[2].cells[1].text = "Cộng Phần A (sau giảm giá)"
    tbl.rows[2].cells[2].text = "27.272.724"

    tbl.rows[3].cells[1].text = "Thuế GTGT 10%"
    tbl.rows[3].cells[2].text = "2.727.272"

    tbl.rows[4].cells[1].text = "TỔNG GIÁ TRỊ HỢP ĐỒNG (A + B)"
    tbl.rows[4].cells[2].text = "29.999.996"

    # Call prevention
    eng._prevent_table_row_splits(doc)

    # All rows must have <w:cantSplit/>
    for ri, row in enumerate(tbl.rows):
        tr = row._tr
        trpr = tr.find(qn("w:trPr"))
        assert trpr is not None, f"Row {ri} missing trPr"
        assert trpr.find(qn("w:cantSplit")) is not None, f"Row {ri} missing cantSplit"

    # Summary rows (rows 2, 3) must have <w:keepNext/> on their paragraphs to prevent breaking before the last total row
    for ri in (2, 3):
        tr = tbl.rows[ri]._tr
        paras = list(tr.iter(qn("w:p")))
        assert len(paras) > 0
        has_keep_next = any(p.find(qn("w:pPr")) is not None and p.find(qn("w:pPr")).find(qn("w:keepNext")) is not None for p in paras)
        assert has_keep_next, f"Row {ri} should have keepNext"
