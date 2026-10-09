"""3 bug báo cáo tiếp theo trên AI Contract Copilot (round 2):
1. Precheck báo thiếu dù CRM đã có / "Bổ sung thông tin" không lưu được - nguyên nhân thật: _apply_legal_overrides() tạo
   TRÙNG Contact mỗi lần precheck tự chạy lại (focus tab) vì không dedup + FE không đồng bộ lại contact_id vừa tạo; và
   update_customer() (crm_customer_service) NULL hoá email_normalized/phone_normalized/source khi patch chỉ sửa 1 trường
   không liên quan (bug CRM gốc, không chỉ ảnh hưởng Copilot).
2. AI Legal Check báo sai giá trị - nguyên nhân thật: docx_to_clauses() làm LỆCH CỘT khi 1 ô giữa bảng rỗng (lọc bỏ ô
   trống rồi mới nối "|"), khiến AI đọc nhầm đơn giá thành SL. Thêm kiểm tra tài chính bằng Decimal độc lập với AI.
3. Lịch thanh toán sai - nguyên nhân thật: AI không hề thấy quote.data.paymentPlan (báo giá có lịch riêng, vd 40/60) nên
   luôn rơi về gợi ý mặc định 50/30/20 trong _STANDARD_TERMS."""
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent))

from app.modules.all_platform.services import contract_ai_service as ai
from app.modules.all_platform.services import contract_source_service as src
from app.modules.all_platform.services.contract_docx_engine import docx_to_clauses


# ───────── 1a. update_customer(): patch 1 trường không được null hoá trường khác ─────────
def test_normalize_payload_does_not_null_unrelated_fields_on_partial_patch():
    from app.modules.all_platform.services.crm_customer_service import _normalize_payload

    out = _normalize_payload({"tax_code": "0312345678"})
    assert out == {"tax_code": "0312345678"}             # KHÔNG tự thêm email_normalized/phone_normalized/source = None
    assert "email_normalized" not in out and "phone_normalized" not in out


def test_normalize_payload_still_normalizes_when_email_or_phone_actually_patched():
    from app.modules.all_platform.services.crm_customer_service import _normalize_payload

    out = _normalize_payload({"email": "A@B.COM", "phone": "0901234567"})
    assert out["email_normalized"] and out["phone_normalized"]   # vẫn chuẩn hoá khi field đó THẬT SỰ có trong patch


# ───────── 1b. _apply_legal_overrides(): không tạo trùng Contact ─────────
def test_apply_legal_overrides_reuses_existing_contact_instead_of_duplicating(monkeypatch):
    from app.modules.all_platform.routers import contract as cr

    calls = {"create": 0, "update": 0}
    monkeypatch.setattr("app.modules.all_platform.services.crm_customer_service.update_customer", lambda *a, **k: None)
    monkeypatch.setattr("app.modules.all_platform.services.crm_contact_service.find_duplicate_contacts",
                         lambda customer_id, user, phone=None, email=None: [{"id": "CT-EXIST", "same_customer": True}])

    def fake_create(*a, **k):
        calls["create"] += 1
        return {"id": "CT-NEW"}

    def fake_update(*a, **k):
        calls["update"] += 1
        return {"id": "CT-EXIST"}

    monkeypatch.setattr("app.modules.all_platform.services.crm_contact_service.create_contact", fake_create)
    monkeypatch.setattr("app.modules.all_platform.services.crm_contact_service.update_contact", fake_update)

    deal = {"customer_id": "C1"}
    overrides = type("O", (), {"model_dump": lambda self, exclude_none=True: {"contact_name": "Thảo", "contact_phone": "090"}})()
    out, contact_id, warnings = cr._apply_legal_overrides(deal, overrides, None, {"id": "u1"}, save_to_crm=True)
    assert calls == {"create": 0, "update": 1}     # KHÔNG tạo mới - tìm thấy trùng (phone) thì update đúng người đã có
    assert contact_id == "CT-EXIST"


def test_apply_legal_overrides_creates_when_truly_no_duplicate(monkeypatch):
    from app.modules.all_platform.routers import contract as cr

    monkeypatch.setattr("app.modules.all_platform.services.crm_customer_service.update_customer", lambda *a, **k: None)
    monkeypatch.setattr("app.modules.all_platform.services.crm_contact_service.find_duplicate_contacts",
                         lambda customer_id, user, phone=None, email=None: [])
    created = {}
    monkeypatch.setattr("app.modules.all_platform.services.crm_contact_service.create_contact",
                         lambda customer_id, patch, user: created.update(id="CT-NEW") or {"id": "CT-NEW"})
    deal = {"customer_id": "C1"}
    overrides = type("O", (), {"model_dump": lambda self, exclude_none=True: {"contact_name": "Ai đó mới", "contact_phone": "099"}})()
    out, contact_id, warnings = cr._apply_legal_overrides(deal, overrides, None, {"id": "u1"}, save_to_crm=True)
    assert contact_id == "CT-NEW" and created.get("id") == "CT-NEW"


# ───────── 2a. docx_to_clauses(): không lệch cột khi 1 ô giữa bảng rỗng ─────────
def test_docx_to_clauses_keeps_column_alignment_when_middle_cell_empty():
    import io
    from docx import Document

    doc = Document()
    doc.add_paragraph("ĐIỀU 2. PHẠM VI CÔNG VIỆC")
    table = doc.add_table(rows=2, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Hạng mục", "Mô tả", "SL", "Đơn giá"
    row = table.rows[1].cells
    row[0].text, row[1].text, row[2].text, row[3].text = "test", "", "1", "159"   # ô Mô tả RỖNG (giống bug thật)
    buf = io.BytesIO()
    doc.save(buf)
    clauses = docx_to_clauses(buf.getvalue())
    body = next(c["body"] for c in clauses if "ĐIỀU 2" in c["title"])
    assert "test | — | 1 | 159" in body     # giữ đúng 4 cột (ô rỗng = "—"), KHÔNG dồn "1" lên vị trí cột Mô tả
    assert "test | 1 | 159" not in body     # dạng lệch cột (bug cũ) không còn xảy ra


def test_docx_to_clauses_skips_fully_blank_rows():
    import io
    from docx import Document

    doc = Document()
    doc.add_paragraph("ĐIỀU 2. X")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "a", "b"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "", ""
    buf = io.BytesIO()
    doc.save(buf)
    clauses = docx_to_clauses(buf.getvalue())
    body = next(c["body"] for c in clauses if "ĐIỀU 2" in c["title"])
    assert body.count("\n") == 0 and "a | b" in body   # dòng trống hoàn toàn vẫn bị bỏ, không chèn "— | —"


# ───────── 2b. Kiểm tra tài chính bằng Decimal, KHÔNG qua AI ─────────
def test_verify_contract_financials_flags_real_mismatch_not_ai_opinion():
    quote = {"totalAmount": 111375177.0, "currency": "VND", "quoteNumber": "202610090150"}
    out = ai._verify_contract_financials(quote, 100000000.0)
    assert len(out) == 1 and out[0]["severity"] == "error"
    assert out[0]["quoteValue"] == 111375177.0 and out[0]["contractValue"] == 100000000.0
    assert round(out[0]["delta"]) == 100000000 - 111375177


def test_verify_contract_financials_silent_when_matching_within_rounding():
    quote = {"totalAmount": 111375177.0, "currency": "VND"}
    assert ai._verify_contract_financials(quote, 111375177.0) == []
    assert ai._verify_contract_financials(quote, 111375177.3) == []    # trong dung sai làm tròn


def test_verify_contract_financials_silent_when_insufficient_data():
    assert ai._verify_contract_financials(None, 100.0) == []
    assert ai._verify_contract_financials({"totalAmount": 100.0}, None) == []


def test_review_risk_score_capped_when_verified_mismatch_even_if_ai_says_safe(monkeypatch):
    async def fake_chat(system, user):
        return {"score": 95, "findings": [{"severity": "ok", "title": "An toàn", "detail": "..."}]}
    monkeypatch.setattr(ai, "_call_chat_json", fake_chat)
    quote = {"totalAmount": 111375177.0, "currency": "VND", "quoteNumber": "Q1"}
    import asyncio
    result = asyncio.run(ai.review_contract_risk([], quote, 50000000.0, None))
    assert result["score"] <= 30       # điểm KHÔNG được mâu thuẫn với lỗi tài chính đã xác nhận, dù AI tự chấm 95
    assert len(result["verifiedFindings"]) == 1 and result["verifiedFindings"][0]["severity"] == "error"


def test_review_risk_keeps_ai_score_when_no_verified_mismatch(monkeypatch):
    async def fake_chat(system, user):
        return {"score": 90, "findings": []}
    monkeypatch.setattr(ai, "_call_chat_json", fake_chat)
    quote = {"totalAmount": 111375177.0, "currency": "VND", "quoteNumber": "Q1"}
    import asyncio
    result = asyncio.run(ai.review_contract_risk([], quote, 111375177.0, None))
    assert result["score"] == 90 and result["verifiedFindings"] == []


# ───────── Legal Check "Xem & xử lý": finding phải mang theo field "clause" (đúng tiêu đề ĐIỀU n.) để FE mở đúng mục ─────────
def test_review_risk_keeps_clause_field_so_fe_can_jump_to_right_section(monkeypatch):
    async def fake_chat(system, user):
        return {"score": 70, "findings": [
            {"severity": "warn", "title": "Thiếu điều khoản phạt", "detail": "...", "clause": "ĐIỀU 4. TRIỂN KHAI & NGHIỆM THU"},
            {"severity": "ok", "title": "An toàn", "detail": "..."},   # AI quên điền clause - vẫn phải có field, không KeyError
        ]}
    monkeypatch.setattr(ai, "_call_chat_json", fake_chat)
    import asyncio
    result = asyncio.run(ai.review_contract_risk([], None, None, None))
    assert result["findings"][0]["clause"] == "ĐIỀU 4. TRIỂN KHAI & NGHIỆM THU"
    assert result["findings"][1]["clause"] == ""      # thiếu field -> mặc định rỗng, FE không hiện nút "Xem & xử lý"


# ───────── 3. Lịch thanh toán thật của báo giá được đưa vào prompt AI, không chỉ 50/30/20 mặc định ─────────
def test_format_quote_context_includes_real_payment_plan_not_hardcoded_default():
    quote = {
        "quoteNumber": "202610090150", "currency": "VND", "status": "approved",
        "subtotalAmount": 103125177.0, "vatAmount": 8250000.0, "totalAmount": 111375177.0,
        "items": [{"description": "Q-Kiosk", "quantity": 1, "unitPrice": 103125000, "vatRate": 8, "totalAmount": 111375000}],
        "data": {"paymentPlan": [{"phase": "Đợt 1", "percent": 40}, {"phase": "Đợt 2", "percent": 60}]},
    }
    text = ai._format_quote_context(quote)
    assert "40%" in text and "60%" in text
    assert "Lịch thanh toán của báo giá này" in text
    # Chỉ là DỮ LIỆU THAM KHẢO, không phải câu lệnh để AI chép nguyên văn vào điều khoản (xem _DRAFT_SYSTEM_PROMPT cho quy tắc "không copy").
    assert "BẮT BUỘC" not in text


def test_format_quote_context_falls_back_to_standard_terms_hint_when_no_plan():
    quote = {"quoteNumber": "Q2", "currency": "VND", "status": "approved", "totalAmount": 100.0, "items": [], "data": {}}
    text = ai._format_quote_context(quote)
    assert "không có lịch thanh toán riêng" in text


def test_standard_terms_drops_5030_20_hint_when_quote_has_own_plan():
    quote_with_plan = {"data": {"paymentPlan": [{"percent": 40}, {"percent": 60}]}}
    quote_without = {"data": {}}
    assert "50% khi ký" not in ai._standard_terms_for(quote_with_plan)
    assert "50% khi ký" in ai._standard_terms_for(quote_without)
    assert "50% khi ký" in ai._standard_terms_for(None)


def test_payment_plan_mismatch_warning_flags_ai_draft_disagreeing_with_quote():
    quote = {"data": {"paymentPlan": [{"phase": "Đợt 1", "percent": 40}, {"phase": "Đợt 2", "percent": 60}]}}
    clauses = [
        {"title": "ĐIỀU 1", "body": "x"}, {"title": "ĐIỀU 2", "body": "y"},
        {"title": "ĐIỀU 3. GIÁ TRỊ & THANH TOÁN", "body": "Thanh toán 50% khi ký, 30% khi bàn giao, 20% sau nghiệm thu."},
    ]
    warnings = src.payment_plan_mismatch_warning(quote, clauses)
    assert len(warnings) == 1 and "40%/60%" in warnings[0].replace(" ", "") or "40" in warnings[0]


def test_payment_plan_mismatch_warning_silent_when_matching():
    quote = {"data": {"paymentPlan": [{"phase": "Đợt 1", "percent": 40}, {"phase": "Đợt 2", "percent": 60}]}}
    clauses = [{"title": "ĐIỀU 3. GIÁ TRỊ & THANH TOÁN", "body": "Thanh toán 40% khi ký, 60% khi bàn giao."}]
    assert src.payment_plan_mismatch_warning(quote, clauses) == []


def test_payment_plan_mismatch_warning_silent_when_quote_has_no_plan():
    clauses = [{"title": "ĐIỀU 3. GIÁ TRỊ & THANH TOÁN", "body": "Thanh toán 50% khi ký, 50% khi bàn giao."}]
    assert src.payment_plan_mismatch_warning({"data": {}}, clauses) == []
    assert src.payment_plan_mismatch_warning(None, clauses) == []
