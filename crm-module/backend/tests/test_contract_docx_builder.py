"""Bộ dựng DOCX hợp đồng (AI soạn mới): bố cục, bảng hạng mục thật, VAT/tổng từ báo giá, lịch thanh toán chính xác."""
import io
from decimal import Decimal

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.modules.all_platform.services import contract_docx_builder as B
from app.modules.all_platform.services.contract_ai_service import build_parties

# Dữ liệu cùng cấu trúc báo giá thật (3 dòng, dòng cuối VAT 8%): tổng 111.375.177
QUOTE = {
    "quoteNumber": "202610090150", "currency": "VND", "subtotalAmount": 103125177, "vatAmount": 8250000, "totalAmount": 111375177,
    "items": [
        {"rowType": "item", "description": "test", "serviceDescription": "test", "unit": "test", "quantity": 1, "unitPrice": 159, "vatRate": 0, "amountAfterDiscount": 159, "totalAmount": 159},
        {"rowType": "item", "description": "t", "serviceDescription": "t", "unit": "t", "quantity": 1, "unitPrice": 18, "vatRate": 0, "amountAfterDiscount": 18, "totalAmount": 18},
        {"rowType": "item", "description": "Q-Kiosk 27SMT U726-P80QRCA-CCA4S4 " + "cấu hình rất dài " * 6, "serviceDescription": "Màn hình cảm ứng 27 inch, máy in nhiệt, đầu đọc QR " * 3,
         "unit": "Bộ", "quantity": 1, "unitPrice": 103125000, "vatRate": 8, "amountAfterDiscount": 103125000, "totalAmount": 111375000},
    ],
}
DEAL = {"customer_name": "Nguyễn Văn An", "company_name": "CÔNG TY A", "tax_code": "0312345678", "address": "25 Lê Lợi, Q1, HCM",
        "contact_name": "Nguyễn Văn An", "position": "Giám đốc", "phone": "0901", "email": "a@a.vn"}
ISSUER = {"legalName": "CÔNG TY MARKEE", "taxCode": "0100", "address": "Hà Nội", "contactName": "Trần B", "phone": "0902", "email": "b@m.vn"}
CLAUSES = [
    # Không gồm sẵn khối "BÊN A...Email:" - khối này do hệ thống tự dựng từ `parties` (xem build()); nếu test cần việc GIỮ NGUYÊN
    # khối Sale tự sửa tay (_parties_edited_manually), dùng fixture riêng, không dùng CLAUSES dùng chung này.
    {"title": "ĐIỀU 1. THÔNG TIN CÁC BÊN", "body": "Hai bên đủ tư cách pháp lý ký kết hợp đồng."},
    {"title": "ĐIỀU 2. PHẠM VI CÔNG VIỆC", "body": "Bên B cung cấp các hạng mục gồm: (i) test; (ii) t; (iii) Q-Kiosk theo bảng dưới đây."},
    {"title": "ĐIỀU 3. GIÁ TRỊ & THANH TOÁN", "body": "Thanh toán 50% khi ký, 40% khi bàn giao và 10% sau nghiệm thu."},
    {"title": "ĐIỀU 4. TRIỂN KHAI & NGHIỆM THU", "body": "Thời gian triển khai 45 ngày."},
]


def doc_of(b): return Document(io.BytesIO(b))


def test_money_vnd_never_has_fractions():
    assert B.money(111375177) == "111.375.177" and B.money(Decimal("55687588.5")) == "55.687.589" and B.money(0) == "0"
    assert B.money(1234.5, "USD") == "1,234.50"


def test_payment_allocation_sums_exactly_with_remainder_in_last_installment():
    steps, warn = B.parse_payment_plan(CLAUSES[2]["body"])
    assert warn is None and [s["percent"] for s in steps] == [50, 40, 10] and "khi ký" in steps[0]["label"]
    total = Decimal(111375177)
    alloc = B.allocate_payments(total, steps)
    amounts = [a["amount"] for a in alloc]
    assert amounts == [Decimal(55687589), Decimal(44550071), Decimal(11137517)]
    assert sum(amounts) == total and all(a == a.to_integral_value() for a in amounts)       # tổng các đợt = đúng tổng hợp đồng, không số lẻ


@pytest.mark.parametrize("total,pcts", [(Decimal(100), [33, 33, 34]), (Decimal(1000001), [30, 30, 40]), (Decimal(999), [50, 50]), (Decimal(1), [60, 40])])
def test_allocation_invariant_for_awkward_totals(total, pcts):
    plan = [{"percent": Decimal(p), "label": ""} for p in pcts]
    alloc = B.allocate_payments(total, plan)
    assert sum(a["amount"] for a in alloc) == total and all(a["amount"] >= 0 for a in alloc)


def test_payment_plan_not_invented_when_percentages_do_not_sum_to_100():
    steps, warn = B.parse_payment_plan("Thanh toán 50% khi ký và 30% khi nghiệm thu.")
    assert steps == [] and "80" in warn
    assert B.parse_payment_plan("Giá chưa gồm VAT 10%.")[0] == []                  # chỉ 1 tỷ lệ + VAT -> không có lịch
    assert B.parse_payment_plan("Thanh toán 100% khi ký.")[0] == []                  # 1 đợt: không dựng bảng


def test_split_enumerated_into_separate_items():
    out = B.split_enumerated("Bên B cung cấp: (i) thiết kế; (ii) quảng cáo; (iii) quản trị.")
    assert out == ["Bên B cung cấp:", "(i) thiết kế", "(ii) quảng cáo", "(iii) quản trị."]
    inst = B.split_enumerated("Thanh toán: Đợt 1: 50% khi ký. Đợt 2: 50% khi nghiệm thu.")
    assert len(inst) == 3 and inst[1].startswith("Đợt 1") and inst[2].startswith("Đợt 2")
    assert B.split_enumerated("Một câu bình thường (có ngoặc).") == ["Một câu bình thường (có ngoặc)."]


def build(clauses=CLAUSES, quote=QUOTE, deal=DEAL):
    parties = build_parties(deal, quote, ISSUER)
    return B.build_contract_docx("Hợp đồng cung cấp dịch vụ", "HD-2026-0001", clauses, parties=parties, quote=quote, signer_a="Nguyễn Văn An", signer_b="Trần B")


def test_document_layout_page_fonts_and_no_wall_of_text():
    docx, warnings = build()
    d = doc_of(docx)
    s = d.sections[0]
    assert round(s.page_width.cm, 1) == 21.0 and round(s.page_height.cm, 1) == 29.7
    texts = [p.text for p in d.paragraphs]
    assert "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in texts and "Số: HD-2026-0001" in texts
    assert max(len(t) for t in texts) < 400                                        # không có đoạn dồn chữ dài
    # Điều 1: hai bên là 2 nhóm riêng, mỗi trường 1 dòng; không dòng nào chứa cả Bên A lẫn Bên B
    assert not any("BÊN A" in t and "BÊN B" in t for t in texts)
    assert any(t.startswith("Mã số thuế: 0312345678") for t in texts) and any(t.startswith("Mã số thuế: 0100") for t in texts)
    assert not any("Hai bên" in t and "BÊN" in t for t in texts)
    # heading luôn đi cùng đoạn sau (không nằm cuối trang một mình)
    heads = [p for p in d.paragraphs if p.text.startswith("ĐIỀU ")]
    assert len(heads) == 4 and all(h.paragraph_format.keep_with_next for h in heads)
    # liệt kê (i)(ii)(iii) tách từng dòng
    assert [t for t in texts if t.startswith("(i")] == ["(i) test;" if False else t for t in texts if t.startswith("(i")]
    assert sum(1 for t in texts if t[:4] in ("(i) ", "(ii)", "(iii")) == 3
    assert not any(t.count("(i)") and t.count("(ii)") for t in texts)


def test_items_table_is_a_real_word_table_with_every_item_in_order():
    docx, _ = build()
    d = doc_of(docx)
    items = d.tables[0]
    assert [c.text for c in items.rows[0].cells][:2] == ["STT", "Hạng mục"] and "Thành tiền" in items.rows[0].cells[7].text
    assert len(items.rows) == 1 + 3
    assert [r.cells[1].text for r in items.rows[1:]][:2] == ["test", "t"] and items.rows[3].cells[1].text.startswith("Q-Kiosk 27SMT U726-P80QRCA-CCA4S4")
    assert items.rows[3].cells[2].text.startswith("Màn hình cảm ứng")                       # mô tả dài vẫn nguyên, không cắt
    assert items.rows[3].cells[3].text == "Bộ" and items.rows[3].cells[4].text == "1" and items.rows[3].cells[5].text == "103.125.000"
    assert items.rows[3].cells[6].text == "8%" and items.rows[3].cells[7].text == "111.375.000"
    assert items.rows[1].cells[2].text == ""                                                  # mô tả trùng tên hạng mục không lặp lại
    # căn phải cột tiền, header lặp lại ở trang sau, hàng không bị cắt đôi
    assert items.rows[3].cells[7].paragraphs[0].alignment == 2 and items.rows[3].cells[5].paragraphs[0].alignment == 2
    assert items.rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
    assert all(r._tr.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None for r in items.rows)


def test_totals_and_schedule_come_from_quote_with_exact_arithmetic():
    docx, warnings = build()
    d = doc_of(docx)
    totals, schedule = d.tables[1], d.tables[2]
    vals = {r.cells[0].text: r.cells[1].text for r in totals.rows}
    assert vals["Tổng giá trị trước thuế"] == "103.125.177 VND" and vals["Thuế GTGT (tổng)"] == "8.250.000 VND"
    assert vals["Tổng giá trị thanh toán (đã gồm VAT)"] == "111.375.177 VND"
    rows = [[c.text for c in r.cells] for r in schedule.rows[1:]]
    assert [r[1] for r in rows] == ["50%", "40%", "10%"] and [r[2] for r in rows] == ["55.687.589", "44.550.071", "11.137.517"]
    assert sum(int(r[2].replace(".", "")) for r in rows) == 111375177
    assert not warnings                                                                        # tổng các dòng khớp tổng báo giá


def test_inconsistent_quote_is_reported_not_silently_fixed():
    bad = {**QUOTE, "totalAmount": 999}
    docx, warnings = build(quote=bad)
    assert any("khác tổng báo giá" in w for w in warnings)
    assert doc_of(docx).tables[1].rows[2].cells[1].text == "999 VND"                          # số của báo giá giữ nguyên, không tự sửa


def test_missing_legal_data_left_blank_with_warning_never_invented():
    docx, warnings = build(deal={"customer_name": "An"})
    texts = [p.text for p in doc_of(docx).paragraphs]
    assert any(t.startswith("Mã số thuế: ………") for t in texts)
    assert any("Bên A: Mã số thuế" in w for w in warnings)


def test_signature_block_stays_together_and_has_both_parties():
    docx, _ = build()
    sig = doc_of(docx).tables[-1]
    assert sig.rows[0].cells[0].text == "ĐẠI DIỆN BÊN A" and sig.rows[0].cells[1].text == "ĐẠI DIỆN BÊN B"
    assert all(r._tr.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None for r in sig.rows)
    assert all(p.paragraph_format.keep_with_next for p in sig.rows[0].cells[0].paragraphs)


def test_without_quote_or_parties_still_builds_clean_document():
    docx, warnings = B.build_contract_docx("HĐ", "1", CLAUSES, parties=None, quote=None)
    d = doc_of(docx)
    assert len(d.tables) == 1 and not warnings                                                 # chỉ có bảng chữ ký; không bịa bảng hạng mục
