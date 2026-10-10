"""Dựng DOCX hợp đồng CHUYÊN NGHIỆP cho chế độ "AI soạn mới" (không có mẫu DOCX): A4, Times New Roman, quốc hiệu, điều khoản có heading riêng,
Điều 1 = hai nhóm Bên A/Bên B theo từng trường, Điều 2 = BẢNG WORD THẬT hạng mục lấy từ báo giá CRM, Điều 3 = tổng trước thuế/VAT/thanh toán từ báo giá
+ lịch thanh toán tính bằng số học thập phân (phần dư dồn vào đợt cuối), chữ ký không tách trang. DOCX là nguồn chuẩn; PDF chuyển từ DOCX bằng LibreOffice.

Số tiền/VAT CHỈ lấy từ báo giá; tỷ lệ/thời hạn thanh toán chỉ lấy từ nội dung điều khoản đã được người dùng duyệt (không tự đặt)."""
from __future__ import annotations

import re
from decimal import ROUND_HALF_UP, Decimal
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.modules.all_platform.services.contract_docx_engine import save_document
from app.modules.all_platform.services.supabase_quote_service import flatten_quote_items

FONT = "Times New Roman"
BLANK = "………………………"
C, J, R, L = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.LEFT


# ───────────────────────── tiền tệ ─────────────────────────
def money(value: Any, currency: str = "VND") -> str:
    """VND: số nguyên, dấu chấm ngăn nghìn (không bao giờ có số lẻ). Tiền tệ khác: 2 số lẻ."""
    try:
        d = Decimal(str(value if value is not None else 0))
    except Exception:  # noqa: BLE001
        return ""
    if (currency or "VND").upper() == "VND":
        n = int(d.quantize(Decimal(1), rounding=ROUND_HALF_UP))
        return f"{n:,}".replace(",", ".")
    q = d.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"{q:,.2f}"


def qty_text(value: Any) -> str:
    try:
        d = Decimal(str(value))
    except Exception:  # noqa: BLE001
        return ""
    return f"{int(d)}" if d == d.to_integral_value() else f"{d.normalize():f}".replace(".", ",")


def item_total_incl_vat(item: dict) -> Decimal:
    """Thành tiền (gồm VAT) của 1 hạng mục: totalAmount của báo giá; chỉ khi báo giá không có thì amountAfterDiscount x (1 + VAT%)."""
    ta, base = item.get("totalAmount"), Decimal(str(item.get("amountAfterDiscount") or 0))
    if ta is not None and (Decimal(str(ta)) > 0 or base == 0):
        return Decimal(str(ta))
    return base + base * Decimal(str(item.get("vatRate") or 0)) / Decimal(100)


def quote_totals(quote: dict) -> dict[str, Decimal]:
    return {
        "subtotal": Decimal(str(quote.get("subtotalAmount") or 0)),
        "vat": Decimal(str(quote.get("vatAmount") or 0)),
        "total": Decimal(str(quote.get("totalAmount") or 0)),
    }


def check_quote_consistency(quote: dict) -> list[str]:
    """Báo giá là nguồn chuẩn - KHÔNG sửa số, chỉ cảnh báo nếu tổng các dòng lệch tổng báo giá."""
    items = [i for i in flatten_quote_items(quote.get("items")) if i.get("rowType") != "section"]
    if not items:
        return []
    t = quote_totals(quote)
    s = sum((item_total_incl_vat(i) for i in items), Decimal(0))
    return [f"Tổng các dòng ({money(s, quote.get('currency'))}) khác tổng báo giá ({money(t['total'], quote.get('currency'))}) - hãy kiểm tra báo giá."] if abs(s - t["total"]) > Decimal("0.5") else []


# ───────────────────────── lịch thanh toán ─────────────────────────
_PCT = re.compile(r"(\d{1,3}(?:[.,]\d+)?)\s*%\s*(.*?)(?=\s+(?:và\s+)?\d{1,3}(?:[.,]\d+)?\s*%|[;,.\n]|$)")


def parse_payment_plan(text: str) -> tuple[list[dict], str | None]:
    """Đọc các đợt từ nội dung Điều thanh toán ĐÃ DUYỆT, VD '50% khi ký, 40% khi bàn giao và 10% sau nghiệm thu'. Chỉ trả lịch khi tổng đúng 100%;
    mục liên quan VAT/thuế bị loại. Trả (đợt[], cảnh báo). Không tự đặt tỷ lệ/thời hạn."""
    found = []
    for m in _PCT.finditer(text or ""):
        label = re.sub(r"^(?:và|,|;|\s)+", "", m.group(2).strip())
        if re.search(r"\b(vat|thuế|gtgt)\b", label.lower()) or re.search(r"\b(vat|thuế)\b", (text or "")[max(0, m.start() - 12): m.start()].lower()):
            continue
        found.append({"percent": Decimal(m.group(1).replace(",", ".")), "label": label})
    if len(found) < 2:
        return [], None
    total_pct = sum((f["percent"] for f in found), Decimal(0))
    if total_pct != Decimal(100):
        return [], f"Các tỷ lệ thanh toán tìm thấy cộng lại {total_pct}% (khác 100%) nên không dựng bảng đợt thanh toán - kiểm tra nội dung Điều thanh toán."
    return found, None


def allocate_payments(total: Decimal, plan: list[dict], currency: str = "VND") -> list[dict]:
    """Số tiền từng đợt = tổng x % (làm tròn HALF_UP theo đơn vị nhỏ nhất của tiền tệ); phần dư dồn vào đợt CUỐI để tổng các đợt = đúng tổng hợp đồng."""
    exp = Decimal(1) if (currency or "VND").upper() == "VND" else Decimal("0.01")
    out, running = [], Decimal(0)
    for idx, step in enumerate(plan):
        if idx == len(plan) - 1:
            amount = total - running
        else:
            amount = (total * step["percent"] / Decimal(100)).quantize(exp, rounding=ROUND_HALF_UP)
            running += amount
        out.append({**step, "amount": amount})
    return out


# ───────────────────────── tách liệt kê trong 1 đoạn ─────────────────────────
_ROMAN = re.compile(r"\(\s*(?:i{1,3}|iv|v|vi{1,3}|ix|x)\s*\)", re.I)
_INST = re.compile(r"(?=\b[Đđ]ợt\s*\d+\s*[:\-–(])")


def split_enumerated(text: str) -> list[str]:
    """Đoạn dồn '(i) ...; (ii) ...' hoặc 'Đợt 1: ... Đợt 2: ...' -> 1 dòng dẫn + từng mục riêng. Không đổi nội dung chữ."""
    t = (text or "").strip()
    marks = list(_ROMAN.finditer(t))
    if len(marks) >= 2:
        parts = [t[:marks[0].start()].strip()] + [t[m.start(): (marks[i + 1].start() if i + 1 < len(marks) else len(t))].strip(" ;,") for i, m in enumerate(marks)]
        return [p for p in parts if p]
    inst = [m.start() for m in _INST.finditer(t)]
    if len(inst) >= 2:
        parts = [t[:inst[0]].strip()] + [t[s: (inst[i + 1] if i + 1 < len(inst) else len(t))].strip(" ;,") for i, s in enumerate(inst)]
        return [p for p in parts if p]
    return [t] if t else []


# ───────────────────────── helper docx ─────────────────────────
def _font(run, size=13, bold=False, italic=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold, run.italic = bold, italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)


def _para(doc_or_cell, text="", *, size=13, bold=False, italic=False, align=J, before=0, after=4, indent=None, hanging=None, keep_next=False, keep_together=True, line=1.25):
    p = doc_or_cell.add_paragraph()
    if text:
        _font(p.add_run(text), size, bold, italic)
    pf = p.paragraph_format
    p.alignment = align
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    # đoạn ngắn giữ nguyên khối; đoạn dài (> 350 ký tự) được phép tách giữa 2 trang (widow/orphan control) để không để trống nửa trang
    pf.keep_with_next, pf.keep_together, pf.widow_control = keep_next, keep_together and len(text or "") < 350, True
    if indent is not None:
        pf.left_indent = Cm(indent)
    if hanging is not None:
        pf.first_line_indent = Cm(-hanging)
    return p


def _label_value(doc, label: str, value: str, indent=0.6, keep_next=True):
    p = _para(doc, "", align=L, after=2, indent=indent + 0.0, keep_next=keep_next)
    _font(p.add_run(f"{label}: "), 13, bold=True)
    _font(p.add_run(value or BLANK), 13)
    return p


def _cell_text(cell, text, *, size=10, bold=False, align=L, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.1
    lines = str(text).split("\n")
    for i, ln in enumerate(lines):
        r = p.add_run(ln)
        _font(r, size, bold, color=color)
        if i < len(lines) - 1:
            r.add_break()


def _shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _set_widths(table, widths_cm):
    tblpr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def _row_flags(row, *, header=False, cant_split=True):
    trpr = row._tr.get_or_add_trPr()
    if header:
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trpr.append(el)
    if cant_split:
        el = OxmlElement("w:cantSplit")
        el.set(qn("w:val"), "true")
        trpr.append(el)


def _keep_row_with_next(row):
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.keep_with_next = True


def _table(doc, rows, cols, widths):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _set_widths(t, widths)
    return t


def _page_field(paragraph, instr):
    run = paragraph.add_run()
    _font(run, 10)
    for kind, text in (("begin", None), (None, f" {instr} "), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            el.text = text
        run._element.append(el)


# ───────────────────────── khối nội dung ─────────────────────────
def _party_block(doc, heading: str, p: dict, who: str, missing: list[str]):
    _para(doc, heading, bold=True, align=L, before=6, after=3, keep_next=True)
    rows = (("Tên đơn vị/cá nhân", p.get("name")), ("Mã số thuế", p.get("tax_code")), ("Địa chỉ", p.get("address")),
            ("Người đại diện", p.get("rep")), ("Chức vụ", p.get("position")), ("Điện thoại", p.get("phone")), ("Email", p.get("email")))
    for ri_, (label, value) in enumerate(rows):
        if not str(value or "").strip() and label in ("Tên đơn vị/cá nhân", "Mã số thuế", "Địa chỉ", "Người đại diện"):   # SĐT/email/chức vụ: để trống được, không cảnh báo
            missing.append(f"{who}: {label}")
        _label_value(doc, label, str(value or "").strip(), keep_next=ri_ < len(rows) - 1)   # cả khối 1 bên đi cùng nhau, nhưng KHÔNG nối dây chuyền sang khối sau


def _items_table(doc, quote: dict, currency: str):
    items = flatten_quote_items(quote.get("items"))
    headers = ["STT", "Hạng mục", "Tính năng / mô tả", "ĐVT", "SL", f"Đơn giá ({currency})", "VAT", f"Thành tiền gồm VAT ({currency})"]
    widths = [1.1, 3.2, 3.4, 1.2, 1.0, 2.3, 1.3, 2.5]
    body = [i for i in items]
    t = _table(doc, 1 + len(body), 8, widths)
    for ci, h in enumerate(headers):
        _cell_text(t.cell(0, ci), h, size=9.5, bold=True, align=C)
        _shade(t.cell(0, ci), "D9E2F3")
    _row_flags(t.rows[0], header=True)
    _keep_row_with_next(t.rows[0])
    n = 0
    for ri, item in enumerate(body, start=1):
        row = t.rows[ri]
        _row_flags(row)
        if ri <= 2:                                   # tiêu đề + 2 dòng đầu đi cùng nhau: bảng không bắt đầu bằng 1 dòng lẻ ở cuối trang
            _keep_row_with_next(row)
        if item.get("rowType") == "section":
            m = row.cells[0].merge(row.cells[-1])
            _cell_text(m, item.get("description") or "", size=10, bold=True)
            _shade(m, "F2F2F2")
            continue
        n += 1
        desc = str(item.get("description") or "").strip()
        feat = str(item.get("serviceDescription") or "").strip()
        vat = item.get("vatRate")
        vals = [str(n), desc, "" if feat == desc else feat, str(item.get("unit") or ""), qty_text(item.get("quantity")),
                money(item.get("unitPrice"), currency) if item.get("unitPrice") is not None else "", f"{qty_text(vat)}%" if vat is not None else "",
                money(item_total_incl_vat(item), currency)]
        aligns = [C, L, L, C, C, R, C, R]
        for ci, (v, a) in enumerate(zip(vals, aligns)):
            _cell_text(row.cells[ci], v, size=10, align=a)
    return t


def _totals_table(doc, quote: dict, currency: str):
    t = quote_totals(quote)
    rows = (("Tổng giá trị trước thuế", t["subtotal"], False), ("Thuế GTGT (tổng)", t["vat"], False), ("Tổng giá trị thanh toán (đã gồm VAT)", t["total"], True))
    tb = _table(doc, len(rows), 2, [11.0, 5.0])
    for ri, (label, value, strong) in enumerate(rows):
        _cell_text(tb.cell(ri, 0), label, size=11, bold=strong)
        _cell_text(tb.cell(ri, 1), f"{money(value, currency)} {currency}", size=11, bold=strong, align=R)
        _row_flags(tb.rows[ri])
        if strong:
            _shade(tb.cell(ri, 0), "F2F2F2"); _shade(tb.cell(ri, 1), "F2F2F2")
    return tb


def _schedule_table(doc, steps: list[dict], currency: str):
    tb = _table(doc, 1 + len(steps), 4, [1.6, 2.0, 3.8, 8.6])
    for ci, h in enumerate(("Đợt", "Tỷ lệ", f"Số tiền ({currency})", "Thời hạn / điều kiện thanh toán")):
        _cell_text(tb.cell(0, ci), h, size=10, bold=True, align=C)
        _shade(tb.cell(0, ci), "D9E2F3")
    _row_flags(tb.rows[0], header=True)
    for ri, s in enumerate(steps, start=1):
        _row_flags(tb.rows[ri])
        _cell_text(tb.cell(ri, 0), str(ri), size=10, align=C)
        _cell_text(tb.cell(ri, 1), f"{qty_text(s['percent'])}%", size=10, align=C)
        _cell_text(tb.cell(ri, 2), money(s["amount"], currency), size=10, align=R)
        _cell_text(tb.cell(ri, 3), s["label"] or "", size=10)


def _clause_body(doc, body: str, first_before: int = 0):
    first = True
    for raw in str(body or "").split("\n"):
        line = raw.strip()
        if not line:
            continue
        pieces = split_enumerated(line)
        for idx, piece in enumerate(pieces):
            is_item = len(pieces) > 1 and idx > 0
            _para(doc, piece, align=J, before=first_before if first else 0, after=3, indent=0.9 if is_item else None, hanging=0.6 if is_item else None)
            first = False


def _strip_parties_block(body: str) -> str:
    """Khối thông tin hai bên do hệ thống chèn đầu Điều 1 (xem _attach_parties): bỏ ra, vì Điều 1 được dựng lại có cấu trúc từ dữ liệu CRM."""
    body = body or ""
    if body.startswith("BÊN A") and "BÊN B" in body:
        end = body.find("Email:", body.find("BÊN B"))
        if end != -1:
            e2 = body.find("\n", end)
            return body[(len(body) if e2 == -1 else e2):].strip()
    return body.strip()


def _extract_parties_block(body: str) -> str:
    """Phần ngược lại của _strip_parties_block() - lấy ĐÚNG khối 'BÊN A...Email:' (không lấy phần văn xuôi phía sau), để so sánh
    với khối hệ thống sẽ tự dựng từ dữ liệu CRM hiện tại (_parties_edited_manually)."""
    body = body or ""
    if body.startswith("BÊN A") and "BÊN B" in body:
        end = body.find("Email:", body.find("BÊN B"))
        if end != -1:
            e2 = body.find("\n", end)
            return body[:(len(body) if e2 == -1 else e2)].strip()
    return ""


def _parties_edited_manually(body: str, parties: dict) -> bool:
    """True nếu Sale đã tự sửa TRỰC TIẾP khối 'BÊN A/BÊN B' trong textarea Điều 1 (khác với khối hệ thống sẽ tự dựng từ dữ liệu
    CRM/đơn vị phát hành hiện tại) - lúc đó PHẢI giữ nguyên đúng chữ Sale đã viết, không tự ghi đè bằng dữ liệu CRM mới nhất mỗi
    lần dựng lại tài liệu (bug cũ: gõ sửa tên/địa chỉ vào Điều 1 rồi lưu, bản DOCX vẫn hiện dữ liệu CRM cũ -> tưởng 'Lưu không
    hoạt động'). Không tự ghi dữ liệu đã sửa ngược lại CRM - chỉ ảnh hưởng tới văn bản hợp đồng này."""
    from app.modules.all_platform.services.contract_ai_service import format_parties_block

    current = _extract_parties_block(body)
    if not current:
        return False
    return current != format_parties_block(parties).strip()


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").upper())


# ───────────────────────── hàm chính ─────────────────────────
def build_contract_docx(
    title: str, contract_number: str, clauses: list[dict], *, parties: dict | None = None, quote: dict | None = None,
    signer_a: str = "", signer_b: str = "", roles: tuple[str, str] | None = None,
) -> tuple[bytes, list[str]]:
    warnings: list[str] = []
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(2), Cm(2.2), Cm(2.2)
    sec.header_distance = sec.footer_distance = Cm(1.1)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(13)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)

    for style_name in ("Header", "Footer"):          # chữ số trang (field) lấy cỡ chữ từ style, không phải từ run
        try:
            doc.styles[style_name].font.size = Pt(10)
            doc.styles[style_name].font.name = FONT
        except KeyError:
            pass
    # header/footer
    hp = sec.header.paragraphs[0]
    hp.alignment = R
    _font(hp.add_run(f"Hợp đồng số {contract_number}" if contract_number else "Hợp đồng"), 9, italic=True, color="666666")
    fp = sec.footer.paragraphs[0]
    fp.alignment = C
    _font(fp.add_run("Trang "), 10)
    _page_field(fp, "PAGE")
    _font(fp.add_run(" / "), 10)
    _page_field(fp, "NUMPAGES")

    currency = ((quote or {}).get("currency") or "VND").upper()
    _para(doc, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, align=C, after=0)
    _para(doc, "Độc lập - Tự do - Hạnh phúc", bold=True, align=C, after=2)
    _para(doc, "—————", align=C, after=10)
    _para(doc, (title or "HỢP ĐỒNG").upper(), size=15, bold=True, align=C, after=2, keep_next=True)
    if contract_number:
        _para(doc, f"Số: {contract_number}", align=C, after=14, keep_next=True)

    missing: list[str] = []
    for idx, clause in enumerate(clauses or []):
        ctitle = str(clause.get("title") or "").strip()
        body = str(clause.get("body") or "")
        key = _norm(ctitle)
        _para(doc, ctitle, bold=True, align=L, before=12, after=4, keep_next=True)
        if idx == 0 and parties and _parties_edited_manually(body, parties):
            # Sale đã tự sửa trực tiếp khối Bên A/B trong textarea Điều 1 - giữ NGUYÊN VĂN, không tự dựng lại từ CRM.
            _clause_body(doc, body, first_before=8)
        elif idx == 0 and parties:
            role_a, role_b = roles or ("Bên sử dụng dịch vụ", "Bên cung cấp dịch vụ")
            _party_block(doc, f"BÊN A ({role_a})", parties["a"], "Bên A", missing)
            _party_block(doc, f"BÊN B ({role_b})", parties["b"], "Bên B", missing)
            _clause_body(doc, _strip_parties_block(body), first_before=8)
        elif quote and quote.get("items") and ("PHẠM VI" in key or idx == 1):
            _clause_body(doc, body)
            _para(doc, "Bảng hạng mục (theo báo giá " + str(quote.get("quoteNumber") or "") + "):", bold=True, align=L, before=4, after=3, keep_next=True)
            _items_table(doc, quote, currency)
            _para(doc, "", after=2)
            warnings += check_quote_consistency(quote)
        elif quote and ("GIÁ TRỊ" in key or "THANH TOÁN" in key or idx == 2):
            _totals_table(doc, quote, currency)
            _para(doc, "", after=2)
            _clause_body(doc, body)
            steps, warn = parse_payment_plan(body)
            if warn:
                warnings.append(warn)
            if steps:
                total = quote_totals(quote)["total"]
                alloc = allocate_payments(total, steps, currency)
                _para(doc, "Lịch thanh toán (số tiền tính từ tổng thanh toán của báo giá):", bold=True, align=L, before=4, after=3, keep_next=True)
                _schedule_table(doc, alloc, currency)
                _para(doc, "", after=2)
        else:
            _clause_body(doc, body)

    if missing:
        warnings.append("Thông tin pháp lý còn thiếu (để trống ………, không tự điền): " + "; ".join(missing))

    _para(doc, "Hợp đồng được lập thành 02 bản có giá trị pháp lý như nhau, mỗi bên giữ 01 bản.", before=10, after=10, keep_next=True)
    sig = _table(doc, 2, 2, [8.0, 8.0])
    for ci, (who, name) in enumerate((("ĐẠI DIỆN BÊN A", signer_a), ("ĐẠI DIỆN BÊN B", signer_b))):
        top, bottom = sig.cell(0, ci), sig.cell(1, ci)
        _cell_text(top, who, size=13, bold=True, align=C)
        _cell_text(bottom, "(Ký, ghi rõ họ tên, đóng dấu)\n\n\n\n" + (name or ""), size=11, align=C)
    for tbl in (sig,):
        tblpr = tbl._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "nil")
            borders.append(el)
        tblpr.append(borders)
    for row in sig.rows:
        _row_flags(row)
    _keep_row_with_next(sig.rows[0])
    return save_document(doc), warnings
